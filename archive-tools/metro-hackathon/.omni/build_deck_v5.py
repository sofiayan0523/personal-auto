# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

INK   = RGBColor(0x17,0x23,0x1C)
TEAL  = RGBColor(0x0B,0x7A,0x75)
DEEP  = RGBColor(0x08,0x3C,0x37)
LEAF  = RGBColor(0x38,0xB8,0x6F)
MINT  = RGBColor(0xC8,0xF3,0xDF)
SUN   = RGBColor(0xFF,0xD3,0x4D)
COIN  = RGBColor(0xF5,0xA5,0x24)
CORAL = RGBColor(0xFF,0x6A,0x3D)
CREAM = RGBColor(0xFF,0xF7,0xDF)
PAPER = RGBColor(0xFF,0xFD,0xF2)
GREY  = RGBColor(0x6E,0x77,0x6F)
WHITE = RGBColor(0xFF,0xFF,0xFF)
FONT  = '微軟正黑體'

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
A = 'personal-auto/final-deck/assets'
V3 = 'personal-auto/final-deck/assets/v3'

def slide_bg(s, color=CREAM):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color

def set_text(tf, lines, default_size=14, default_color=INK, align=PP_ALIGN.LEFT, bold=False, font=FONT, line_spacing=1.12):
    tf.word_wrap = True
    first = True
    for line in lines:
        if isinstance(line, str): line = {'t': line}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = line.get('align', align)
        p.line_spacing = line.get('ls', line_spacing)
        p.space_after = Pt(line.get('sa', 4))
        r = p.add_run(); r.text = line['t']
        r.font.name = line.get('font', font)
        r.font.size = Pt(line.get('size', default_size))
        r.font.bold = line.get('bold', bold)
        r.font.color.rgb = line.get('color', default_color)
        rPr = r._r.get_or_add_rPr()
        ea = rPr.find(qn('a:ea'))
        if ea is None:
            ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
        ea.set('typeface', line.get('font', font))

def card(s, x, y, w, h, fill=PAPER, line_color=INK, line_w=2.5, shadow=0.09, radius=0.12, shadow_color=INK):
    if shadow:
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x+shadow), Inches(y+shadow), Inches(w), Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = shadow_color
        sh.line.fill.background(); sh.shadow.inherit = False
        try: sh.adjustments[0] = radius
        except Exception: pass
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.fill.solid(); c.fill.fore_color.rgb = fill
    c.line.color.rgb = line_color; c.line.width = Pt(line_w)
    c.shadow.inherit = False
    try: c.adjustments[0] = radius
    except Exception: pass
    return c

def chip(s, x, y, w, h, text, fill=SUN, tcolor=INK, size=12, rot=0, bold=True, line_w=2.0):
    c = card(s, x, y, w, h, fill=fill, line_w=line_w, shadow=0.05, radius=0.5)
    if rot: c.rotation = rot
    tf = c.text_frame; tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(tf, [{'t': text, 'size': size, 'bold': bold, 'color': tcolor, 'align': PP_ALIGN.CENTER, 'sa': 0}])
    return c

def txt(s, x, y, w, h, lines, **kw):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(tb.text_frame, lines, **kw)
    return tb

def header(s, num, title, tag=None):
    bar = card(s, 0.55, 0.42, 0.62, 0.62, fill=SUN, line_w=2.5, shadow=0.07, radius=0.25)
    tfb = bar.text_frame; tfb.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(tfb, [{'t': str(num), 'size': 22, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0, 'font': 'Arial Black'}])
    txt(s, 1.4, 0.42, 9.6, 0.8, [{'t': title, 'size': 27, 'bold': True}])
    if tag:
        chip(s, 10.85, 0.5, 2.0, 0.5, tag, fill=MINT, size=12)

def footer(s, page):
    txt(s, 0.55, 7.08, 9.0, 0.35, [{'t': '從碳客變捷客 · 2026 捷運盃黑客松決賽簡報 (提案審閱完整版) · 想領獎金的捷客', 'size': 9, 'color': GREY}])
    txt(s, 12.2, 7.08, 0.7, 0.35, [{'t': str(page), 'size': 10, 'color': GREY, 'align': PP_ALIGN.RIGHT, 'font': 'Arial'}])

# ============ SLIDE 1 — COVER ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
card(s, 0.55, 0.55, 7.4, 6.4, fill=TEAL, line_w=3, shadow=0.12, radius=0.06)
txt(s, 1.1, 1.0, 6.4, 0.5, [{'t': '2026 捷運盃黑客松 · 決賽提案 (提案審閱對齊版)', 'size': 15, 'bold': True, 'color': MINT}])
txt(s, 1.05, 1.5, 6.6, 2.6, [
    {'t': '從碳客', 'size': 76, 'bold': True, 'color': WHITE, 'sa': 0, 'ls': 1.0},
    {'t': '變捷客', 'size': 76, 'bold': True, 'color': SUN, 'ls': 1.0},
])
hook = card(s, 1.1, 4.55, 5.9, 0.78, fill=SUN, line_w=2.5, shadow=0.08, radius=0.3)
hook.rotation = -1.5
hook.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(hook.text_frame, [{'t': '每一次通勤，都是可驗證的減碳資產', 'size': 19, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0}])
txt(s, 1.1, 5.7, 6.2, 1.0, [
    {'t': '團隊：想領獎金的捷客', 'size': 15, 'bold': True, 'color': WHITE, 'sa': 2},
    {'t': '把捷運點變成「健康共享生活圈」的入口 · 串接市府健康政策與站內招商的永續商模', 'size': 11, 'color': MINT},
])
mc = card(s, 8.35, 1.15, 4.4, 3.5, fill=PAPER, line_w=3, shadow=0.12, radius=0.08)
s.shapes.add_picture(f'{A}/mascot.png', Inches(8.55), Inches(1.4), width=Inches(4.0))
chip(s, 8.6, 4.35, 1.7, 0.5, '決賽 10 強', fill=CORAL, tcolor=WHITE, size=13, rot=-2)
chip(s, 10.45, 4.35, 2.2, 0.5, '互動 Demo 可玩', fill=MINT, size=13, rot=1.5)
txt(s, 8.35, 5.35, 4.5, 1.2, [
    {'t': '線上互動原型 (C端/B2B聯動)', 'size': 12, 'bold': True, 'color': GREY, 'sa': 2},
    {'t': 'from-carbon-metro-vote.pages.dev/demo/', 'size': 13, 'bold': True, 'color': TEAL, 'font': 'Consolas'},
])
footer(s, 1)

# ============ SLIDE 2 — 痛點 (痛點：Go App 資訊化，市民無感、點數冷感) ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 2, '痛點重構：Go App 資訊化，市民無感、點數冷感')
stat = card(s, 0.55, 1.45, 12.2, 1.15, fill=DEEP, line_w=3, shadow=0.1, radius=0.1)
stat.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(stat.text_frame, [
    {'t': '臺北捷運 114 年平均日運量 207.4 萬人次 — 每一趟都是尚未被資產化的減碳貢獻', 'size': 19, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER, 'sa': 2},
    {'t': '資料來源：臺北捷運公司公開營運統計', 'size': 10.5, 'color': MINT, 'align': PP_ALIGN.CENTER, 'sa': 0},
])
pains = [
    ('Go App 資訊化', '現行 App 偏向靜態資訊查詢（時刻表、路線圖、進站狀態），缺乏高黏著的日常情境推播，市民日常通勤缺乏開啟 App 的高頻動機與樂趣。', SUN),
    ('點數閒置過期', '由於缺乏消耗場景與感官回饋，捷運點多被遺忘或過期失效，造成使用者點數冷感，北捷亦面臨點數負債與高過期（Breakage）會計壓力。', COIN),
    ('減碳貢獻無感', '市民雖知搭捷運環保，但「我今天這趟通勤省下多少碳、對城市有何貢獻」從未被具體量化，減碳流於口號，難以轉化為實質的行為改變。', MINT),
]
for i, (t, d, cl) in enumerate(pains):
    x = 0.55 + i * 4.18
    c = card(s, x, 3.0, 3.85, 2.3, fill=PAPER, line_w=2.5, shadow=0.09)
    chip(s, x + 0.25, 3.25, 1.7, 0.55, t, fill=cl, size=15)
    txt(s, x + 0.25, 4.0, 3.4, 1.2, [{'t': d, 'size': 12.5, 'ls': 1.25}])
bot = card(s, 0.55, 5.6, 12.2, 1.15, fill=PAPER, line_w=2.5, shadow=0.09)
bot.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(bot.text_frame, [
    {'t': '核心挑戰：如何破除資訊化 App 的冷感？唯有「引入類皮克敏遊戲化引擎」，串接「市政健康政策」與「站內既存商業租賃」，才能打造高黏著的 Earn-to-Spend 減碳生活圈！', 'size': 14.5, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0},
])
footer(s, 2)

# ============ SLIDE 3 — 提案概念 (市政與商業雙重槓桿：捷運健康共享生活圈) ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 3, '提案概念：捷運健康共享生活圈')
left_c = card(s, 0.55, 1.5, 6.0, 4.5, fill=TEAL, line_w=3, shadow=0.1)
txt(s, 0.9, 1.75, 5.3, 4.1, [
    {'t': '市政健康共享生活圈 (E & S)', 'size': 18, 'bold': True, 'color': SUN, 'sa': 8},
    {'t': '以「捷運點」為核心，直接對接台北市既存健康政策：', 'size': 13.5, 'bold': True, 'color': WHITE, 'sa': 8},
    {'t': '· U-Sport (體育局運動中心) 入場 → 累積減碳與健康進度', 'size': 13, 'color': MINT, 'sa': 4},
    {'t': '· U-Walk (衛生局健康萬步挑戰) 步行 → 出站沿線完成任務', 'size': 13, 'color': MINT, 'sa': 4},
    {'t': '· YouBike (交通局接駁) 騎行 → 省碳直接加成賺點', 'size': 13, 'color': MINT, 'sa': 8},
    {'t': '★ 零新建成本：100% 串接市府既存成熟系統，免發空包彈，點數具有真實健康行為支撐！', 'size': 12.5, 'color': WHITE, 'ls': 1.25},
])
right_c = card(s, 6.85, 1.5, 5.9, 4.5, fill=PAPER, line_w=2.5, shadow=0.1)
txt(s, 7.2, 1.75, 5.2, 4.1, [
    {'t': '站內補給 ✕ 商業招商開發 (G & Profit)', 'size': 18, 'bold': True, 'color': TEAL, 'sa': 8},
    {'t': '串接站內既存租賃服務與特約廠商，轉化為高動能導客商機：', 'size': 13.5, 'bold': True, 'sa': 8},
    {'t': '· 補給：rainGo 共享傘、旅電/ChargeSPOT、小樹屋會議亭', 'size': 13, 'color': GREY, 'sa': 4},
    {'t': '· 點券：亞尼克 YTM 蛋糕自助機、全家、黑潮咖啡站內櫃位', 'size': 13, 'color': GREY, 'sa': 8},
    {'t': '★ 招商加值：原本單純的空間租賃合約被附加「Go App 即時推播、捷運點導客券與核銷」的數位行銷方案！租賃點位招商單價與北捷非票務收入大幅倍增！', 'size': 13, 'ls': 1.25},
])
bar = card(s, 0.55, 6.15, 12.2, 0.78, fill=SUN, line_w=2.5, shadow=0.08)
bar.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(bar.text_frame, [{'t': '日常通勤賺點 ✕ 市政健康共享 ✕ 既存商業租賃 ＝ 一個無新建硬體成本、市民、北捷、合作方三贏的永續閉環！', 'size': 14.5, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0}])
footer(s, 3)

# ============ SLIDE 4 — 遊戲化留存：類皮克敏的遊戲化情感引擎 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 4, '情感引擎：類皮克敏的遊戲化世界觀', tag='內容創新')
shots = [
    (f'{V3}/story-modal-1.png', '開場故事：拯救碳森林危機', 2.55, 2.56),
    (f'{V3}/phone-hatch-modal.png', '健康任務充能 → 夥伴孵化出世', 1.90, 3.94),
    (f'{V3}/phone-forest.png', '散步小隊棲地 ✕ 培育神木升級', 1.90, 3.94),
]
xs = [0.55, 3.45, 5.7]
for (img, cap, w, h), x in zip(shots, xs):
    ytop = 1.55 + (4.1 - h) / 2
    card(s, x, ytop - 0.12, w + 0.24, h + 0.24, fill=PAPER, line_w=2.5, shadow=0.09, radius=0.07)
    s.shapes.add_picture(img, Inches(x + 0.12), Inches(ytop), width=Inches(w))
    txt(s, x - 0.05, 5.85, w + 0.35, 0.6, [{'t': cap, 'size': 11, 'bold': True, 'align': PP_ALIGN.CENTER}])
tier_panel = card(s, 8.0, 1.55, 4.8, 4.35, fill=TEAL, line_w=3, shadow=0.1)
txt(s, 8.35, 1.78, 4.15, 4.0, [
    {'t': '以情感驅動，破除點數冷感', 'size': 17, 'bold': True, 'color': SUN, 'sa': 8},
    {'t': '將乾燥的減碳數字，化為有生命的陪伴：', 'size': 13, 'bold': True, 'color': WHITE, 'sa': 10, 'ls': 1.3},
    {'t': '🌲 碳森林救援故事：', 'size': 13, 'bold': True, 'color': MINT, 'sa': 2},
    {'t': '出站省下的碳，能灌溉瀕危碳森林。小捷情境引導，使命感拉滿。', 'size': 12, 'color': WHITE, 'sa': 7},
    {'t': '🐣 散步夥伴與限定裝扮：', 'size': 13, 'bold': True, 'color': MINT, 'sa': 2},
    {'t': 'YouBike 任務孵化「輪輪」、U-Walk 孵化「步步」、U-Sport 孵化「汗汗」。配備市政府運動汗帶、共享傘等限定裝扮，預防首頁空洞，激發市民搜集魂。', 'size': 12, 'color': WHITE, 'sa': 7},
    {'t': '🌱 每週收成與培育神木：', 'size': 13, 'bold': True, 'color': MINT, 'sa': 2},
    {'t': '週日手動收成結算「種子」，培育森林神木升級，建立高黏著高回訪行為。', 'size': 12, 'color': WHITE},
])
bar = card(s, 8.0, 6.05, 4.8, 0.75, fill=SUN, line_w=2.5, shadow=0.08)
bar.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(bar.text_frame, [{'t': '遊戲化不是點綴，是讓市民從「配合政策」轉為「主動熱愛」的黏著引擎！', 'size': 13, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0}])
footer(s, 4)

# ============ SLIDE 5 — 4 User Scenarios (四個示範場景：在出站時刻，觸發即時補給) ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 5, '四個示範情境：出站最需要的時刻觸發補給')
banner = card(s,Stat_x:=0.55, 1.45, 12.2, 0.78, fill=DEEP, line_w=3, shadow=0.08)
banner.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(banner.text_frame, [{'t': '示範流程聚焦「站點情境 → 推播 → 任務 → 回饋 → 衡量」的完整鏈條，完美對接市民出站生活需求。', 'size': 14.5, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER, 'sa': 0}])
scenarios = [
    ('場景 A：下班健康任務', '出站時，Go App 推播：「今日離萬步還差 2,000 步！步行到中山運動中心入場，即得 15 捷運點，且 U-Walk 蛋充能 100% 孵出步步！」\n→ 資料：出站點、時間、opt-in 健康偏好。', MINT),
    ('場景 B：雨天共享接力', '出站遇雨，App 推播：「本站有 rainGo 共享傘機台；完成甲借乙還，可得 5 捷運點與共享傘徽章！」\n→ 資料：即時天氣、站點、共享傘即時狀態。', SUN),
    ('場景 C：週末低碳生活', '週末減碳行程：「捷運到大安森林公園站 + YouBike 接駁 10 分鐘 + 亞尼克 YTM 蛋糕補給，領取乘數加成。」\n→ 資料：週末模式、YouBike 站點、YTM 銷點店家。', COIN),
    ('場景 D：轉乘空檔服務', '轉乘有 35 分鐘空檔：「中山站有小樹屋電話亭可預約會議，附近快閃櫃或快剪店可用捷運點換券，低等待。」\n→ 資料：站點租賃服務、轉乘時段、會員自選偏好。', MINT),
]
for i, (t, d, cl) in enumerate(scenarios):
    x = 0.55 + i * 3.09
    c = card(s, x, 2.45, 2.95, 4.4, fill=PAPER, line_w=2.5, shadow=0.09)
    chip(s, x + 0.18, 2.7, 2.6, 0.55, t, fill=cl, size=13)
    txt(s, x + 0.18, 3.45, 2.6, 3.25, [{'t': line, 'size': 12, 'ls': 1.25, 'sa': 3} for line in d.split('\n')])
footer(s, 5)

# ============ SLIDE 6 — 3 分鐘 Demo 動線 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 6, '3 分鐘 Demo 動線：一條故事線走完全部機制')
journey = [
    ('① 故事與設定', '3 頁世界觀故事引入\n常用路線 Onboarding 設定\n（深設常用通勤路線）', MINT),
    ('② 出站與結算', '隨車答題互動\n出站：+15點 / +0.23kg\n離峰通勤獲稀有種子', SUN),
    ('③ 任務與孵蛋', '做健康生活圈任務\n（YouBike/U-Sport）\n對應夥伴蛋 100% 孵化', COIN),
    ('④ 森林與收成', '夥伴即時入住森林棲地\n週日手動收成獲種子\n培育神木、升級神木', MINT),
    ('⑤ 花點換券核銷', '100 點兌換全家券\n條碼/QR code 現場核銷\nNumbers 去識別指紋封存', SUN),
]
for i, (t, d, cl) in enumerate(journey):
    x = 0.55 + i * 2.49
    c = card(s, x, 1.6, 2.31, 3.3, fill=PAPER, line_w=2.5, shadow=0.1)
    chip(s, x + 0.18, 1.85, 1.75, 0.55, t, fill=cl, size=13)
    txt(s, x + 0.18, 2.6, 1.99, 2.1, [{'t': line, 'size': 11.5, 'ls': 1.25, 'sa': 3} for line in d.split('\n')])
b = card(s, 0.55, 5.3, 12.2, 1.5, fill=DEEP, line_w=3, shadow=0.1)
txt(s, 0.95, 5.5, 11.5, 1.1, [
    {'t': '評審高光（第 ⑥ 步）：切到 B2B 營運主控台 — 剛在乘客端完成的動作即時滾入 LIVE 互動流！', 'size': 15.5, 'bold': True, 'color': WHITE, 'sa': 3},
    {'t': '真正完成雙端實時聯動原型；模擬「離峰通勤」直接對應北捷「尖峰運能舒緩、離峰乘坐率提升」關鍵營運 KPI。', 'size': 13, 'color': MINT},
])
footer(s, 6)

# ============ SLIDE 7 — Point Mechanism & Reward Loop (點數機制：觸發、判斷、行動、回饋、衡量) ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 7, '點數機制與任務設計：先設計行為，再決定發點')
p_banner = card(s, 0.55, 1.45, 12.2, 0.7, fill=PAPER, line_w=2.5, shadow=0.08)
p_banner.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(p_banner.text_frame, [{'t': '點數機制核心：用捷運點承接「健康、共享、租賃導客」任務，控制成本與資料邊界，實現可預算化管理。', 'size': 14.5, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0}])
flow_box = card(s, 0.55, 2.35, 4.3, 4.5, fill=DEEP, line_w=3, shadow=0.1)
txt(s, 0.85, 2.55, 3.7, 4.1, [
    {'t': '點數運作五步驟', 'size': 17, 'bold': True, 'color': SUN, 'sa': 8},
    {'t': '1. 觸發 (Trigger)：', 'size': 13.5, 'bold': True, 'color': MINT, 'sa': 2},
    {'t': '出站、靠近共享傘或選擇健康挑戰。', 'size': 12.5, 'color': WHITE, 'sa': 6},
    {'t': '2. 判斷 (Judge)：', 'size': 13.5, 'bold': True, 'color': MINT, 'sa': 2},
    {'t': 'App 依站點、時間、天氣、偏好匹配推薦。', 'size': 12.5, 'color': WHITE, 'sa': 6},
    {'t': '3. 行動 (Action)：', 'size': 13.5, 'bold': True, 'color': MINT, 'sa': 2},
    {'t': '步行到場館、騎 YouBike、借傘補電。', 'size': 12.5, 'color': WHITE},
])
details_box = card(s, 5.15, 2.35, 7.6, 4.5, fill=PAPER, line_w=2.5, shadow=0.1)
txt(s, 5.45, 2.55, 7.0, 4.1, [
    {'t': '四大合作任務類型與控制邊界', 'size': 17, 'bold': True, 'color': TEAL, 'sa': 8},
    {'t': '· 健康減碳：搭捷運出站 + 步行到 U-Sport 運動場館或 U-Walk 萬步。', 'size': 13, 'bold': True, 'sa': 2},
    {'t': '  [回饋] 捷運點加成/場館券 [成本控制] 每月名額上限控管 [資料] opt-in 個人授權。', 'size': 12, 'color': GREY, 'sa': 6},
    {'t': '· 共享服務：租借/歸還共享傘 (rainGo)、旅電/ChargeSPOT、騎 YouBike 接駁。', 'size': 13, 'bold': True, 'sa': 2},
    {'t': '  [回饋] 共享任務章/借用折抵券/抽獎券 [成本控制] 由合作廠商分攤，不進行點對點補貼。', 'size': 12, 'color': GREY, 'sa': 6},
    {'t': '· 站內補給：亞尼克 YTM 蛋糕機、全家、黑潮咖啡站內櫃位。', 'size': 13, 'bold': True, 'sa': 2},
    {'t': '  [回饋] 捷運點折抵/特約優惠券 [成本控制] 由商家購買曝光導流包，提供兌換券。', 'size': 12, 'color': GREY, 'sa': 6},
    {'t': '· 空間租賃：小樹屋電話亭、智慧生活站取貨、快閃櫃、快剪。', 'size': 13, 'bold': True, 'sa': 2},
    {'t': '  [回饋] 會員時租券/折扣 [成本控制] 包裹為租賃招商行銷方案，收取導客分成費。', 'size': 12, 'color': GREY},
])
footer(s, 7)

# ============ SLIDE 8 — B2B operations & CPS Live Console (商業導流與 B2B 合作主控台) ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 8, '商業導流：台北捷運 B2B 營運合作夥伴主控台')
card(s, 0.7, 1.5, 3.15, 5.35, fill=PAPER, line_w=3, shadow=0.1, radius=0.07)
s.shapes.add_picture(f'{V3}/phone-b2b-console.png', Inches(1.18), Inches(1.68), height=Inches(4.55))
txt(s, 0.7, 6.3, 3.15, 0.4, [{'t': 'B2B 營運端：LIVE 乘客事件流 ✕ 營運與招商指標', 'size': 11.5, 'bold': True, 'align': PP_ALIGN.CENTER}])
card(s, 4.15, 1.5, 3.15, 5.35, fill=PAPER, line_w=3, shadow=0.1, radius=0.07)
s.shapes.add_picture(f'{V3}/phone-today-settled.png', Inches(4.63), Inches(1.68), height=Inches(4.55))
txt(s, 4.15, 6.3, 3.15, 0.4, [{'t': '乘客端：簡潔、隱私、專注減碳與遊戲化', 'size': 11.5, 'bold': True, 'align': PP_ALIGN.CENTER}])
info = card(s, 7.75, 1.5, 5.05, 5.35, fill=TEAL, line_w=3, shadow=0.1)
txt(s, 8.1, 1.75, 4.4, 4.0, [
    {'t': '可衡量的商模與招商加值方案', 'size': 18, 'bold': True, 'color': SUN, 'sa': 8},
    {'t': '北捷可直接向合作夥伴銷售的四種加值方案：', 'size': 13.5, 'bold': True, 'color': WHITE, 'sa': 6},
    {'t': '1. 既存商家即時推播包 (亞尼克 YTM、旅電、共享傘)', 'size': 12.5, 'color': MINT, 'sa': 4},
    {'t': '2. 既存空間租客導客包 (小樹屋、快閃櫃、快剪)', 'size': 12.5, 'color': MINT, 'sa': 4},
    {'t': '3. 市政健康低碳共創包 (U-Sport 體育局合作)', 'size': 12.5, 'color': MINT, 'sa': 4},
    {'t': '4. 悠遊付生活圈加碼包 (第二階段合作擴充)', 'size': 12.5, 'color': MINT, 'sa': 10},
    {'t': '📊 共同森林看板：', 'size': 13.5, 'bold': True, 'color': WHITE, 'sa': 2},
    {'t': '主控台導入「板南線共同森林」看板，由沿線商家與租客共同捐贈養分，將導客成效具體化為企業 ESG 成效揭露！', 'size': 12, 'color': WHITE, 'ls': 1.25},
])
footer(s, 8)

# ============ SLIDE 9 — Technical Infrastructure: Numbers Trust Ledger (技術硬底子：全場唯一 Numbers 可稽核數據信託) ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 9, '技術硬底子：全場唯一 Numbers 可稽核數據信託')
big = card(s, 0.55, 1.5, 7.6, 3.4, fill=DEEP, line_w=3, shadow=0.1)
txt(s, 0.95, 1.75, 6.9, 3.0, [
    {'t': '發出的每一克減碳、每一點發放，', 'size': 21, 'bold': True, 'color': WHITE, 'sa': 0, 'ls': 1.15},
    {'t': '都經得起外部嚴格審計信託。', 'size': 21, 'bold': True, 'color': SUN, 'sa': 10, 'ls': 1.15},
    {'t': '· 類比雲端發票載具：個人減碳與核銷紀錄，自動封存為去識別化可攜資產，換手機不丟失，第三方可隨時驗證。', 'size': 13.5, 'bold': True, 'color': WHITE, 'sa': 6, 'ls': 1.25},
    {'t': '· 資料指紋 (Provenance)：核銷與減碳憑證自動計算 SHA-256 數據指紋寫入分散式 Ledger，保證數據不可篡改、不重複發放、不灌水。直接對接企業綠色金融或 ESG 審計合規首選！', 'size': 12, 'color': MINT, 'ls': 1.25},
])
tagc = card(s, 0.55, 5.15, 7.6, 0.85, fill=PAPER, line_w=2.5, shadow=0.08)
tagc.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(tagc.text_frame, [{'t': 'Human Truth. Machine Proof. — Numbers Protocol 官方技術支持', 'size': 16, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0, 'font': 'Consolas'}])
ph = card(s, 8.45, 1.5, 3.0, 5.1, fill=PAPER, line_w=3, shadow=0.1, radius=0.07)
s.shapes.add_picture(f'{V3}/phone-voucher-proof.png', Inches(8.9), Inches(1.68), height=Inches(4.35))
txt(s, 8.45, 6.15, 3.0, 0.45, [{'t': '核銷券內嵌「資料指紋已封存」憑證', 'size': 11, 'bold': True, 'align': PP_ALIGN.CENTER}])
pts = card(s, 0.55, 6.15, 7.6, 0.8, fill=MINT, line_w=2.5, shadow=0.07)
pts.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(pts.text_frame, [{'t': '安全與隱私底線：Opt-in 授權、去識別化統計、不儲存個人行蹤、隨時可一鍵關閉', 'size': 13, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0}])
footer(s, 9)

# ============ SLIDE 10 — Financial Projection & Sustainability (財務試算：首年損益與永續會計) ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 10, '可行性與首年財務試算（非空中樓閣）')
nums = [('首年合作金流收入', 'NT$ 2,270 萬', TEAL), ('點數與維運支出', 'NT$ 1,525 萬', COIN), ('可支應營運餘額', 'NT$ 745 萬', LEAF)]
for i, (label, val, cl) in enumerate(nums):
    x = 0.55 + i * 4.18
    c = card(s, x, 1.55, 3.85, 1.7, fill=PAPER, line_w=2.5, shadow=0.1)
    txt(s, x + 0.3, 1.75, 3.3, 0.45, [{'t': label, 'size': 14, 'bold': True, 'color': GREY}])
    txt(s, x + 0.3, 2.2, 3.3, 0.9, [{'t': val, 'size': 28, 'bold': True, 'color': cl, 'font': 'Arial Black'}])
acct = card(s, 0.55, 3.5, 7.9, 3.3, fill=DEEP, line_w=3, shadow=0.1)
txt(s, 0.95, 3.75, 7.2, 2.9, [
    {'t': 'IFRS 15 點數會計與 Breakage 管理（直球對決財務評審）', 'size': 16, 'bold': True, 'color': SUN, 'sa': 8},
    {'t': '· 發點遞延處理：發放時列合約負債（遞延收入），實際核銷或逾期時認列收入。', 'size': 13, 'color': WHITE, 'sa': 5, 'ls': 1.2},
    {'t': '· Breakage 精精算：引入 OPENPOINT 歷史未兌換率（Breakage）估算，不虛增負債壓力。', 'size': 13, 'color': WHITE, 'sa': 5, 'ls': 1.2},
    {'t': '· 成本共同分攤：點數池由站內特約商家、既存租賃商共同出資，北捷年補貼控制在 600 萬內！', 'size': 13, 'color': WHITE, 'sa': 5, 'ls': 1.2},
    {'t': '· 研究底稿：完整分析 LINE POINTS、星巴克與 Marriott 之會計實務。', 'size': 11.5, 'color': MINT},
])
src = card(s, 8.75, 3.5, 4.0, 3.3, fill=PAPER, line_w=2.5, shadow=0.1)
txt(s, 9.05, 3.75, 3.5, 2.9, [
    {'t': '金流與成本細目', 'size': 15, 'bold': True, 'color': TEAL, 'sa': 6},
    {'t': '【收入】', 'size': 12, 'bold': True, 'sa': 2},
    {'t': '· 20 商家推播合作包：1,200 萬 (5萬/月)\n· 任務回饋池商家流：270 萬\n· ESG/品牌共創：500 萬 / 租賃加值：300 萬', 'size': 11.5, 'sa': 4},
    {'t': '【支出】', 'size': 12, 'bold': True, 'sa': 2},
    {'t': '· 回饋履約：225 萬 / 北捷補貼：600 萬\n· 資料治理：200 萬 / 維運客服：500 萬', 'size': 11.5},
])
footer(s, 10)

# ============ SLIDE 11 — 藍圖與誠實標註 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 11, '三階段導入藍圖與風險管理')
phases = [
    ('決賽階段 (現有成果)', '可操作高高擬真原型與 B2B 主控台，100% 離線可用，展示完整商業與健康機制。', MINT),
    ('試辦階段 (半年導入)', '選定 2–3 個示範生活圈站點，對接 U-Sport / Walk 政策，串聯站內既存廠商出站導客。', SUN),
    ('擴大階段 (一年推廣)', '活化常客優惠轉點機制，對接悠遊付百萬商家網絡，引進綠色企業 ESG 數據信託採購。', COIN),
]
for i, (t, d, cl) in enumerate(phases):
    x = 0.55 + i * 4.18
    c = card(s, x, 1.55, 3.85, 2.5, fill=PAPER, line_w=2.5, shadow=0.1)
    chip(s, x + 0.25, 1.75, 2.1, 0.52, t, fill=cl, size=13.5)
    txt(s, x + 0.28, 2.45, 3.3, 1.5, [{'t': d, 'size': 12.5, 'ls': 1.25}])
    if i < 2:
        txt(s, x + 3.85, 2.4, 0.4, 0.8, [{'t': '→', 'size': 18, 'bold': True, 'align': PP_ALIGN.CENTER}])
real = card(s, 0.55, 4.35, 6.0, 2.5, fill=PAPER, line_w=2.5, shadow=0.08)
txt(s, 0.9, 4.55, 5.3, 2.1, [
    {'t': '✓ 需主動管理的風險', 'size': 14.5, 'bold': True, 'color': LEAF, 'sa': 4},
    {'t': '· 跨局處與法人合作：U-Sport、YouBike 與悠遊付需市府行政協調。', 'size': 12, 'sa': 2},
    {'t': '· 隱私防禦：健康與步數資料採最小化，經 opt-in 授權，且隨時可關閉。', 'size': 12, 'sa': 2},
    {'t': '· 廣告審查：租賃空間與推播素材由北捷進行合規審查，不假設租戶可自行販售第三方廣告。', 'size': 12},
])
todo = card(s, 6.75, 4.35, 6.0, 2.5, fill=PAPER, line_w=2.5, shadow=0.08)
txt(s, 7.1, 4.55, 5.3, 2.1, [
    {'t': '△ 北捷審閱重點回應 (誠實標註)', 'size': 14.5, 'bold': True, 'color': CORAL, 'sa': 4},
    {'t': '· 零新建成本：本案完全依託北捷既存的 Go App、Go!Map 與既存租客。', 'size': 12, 'sa': 2},
    {'t': '· 減碳係數 104 / 75 / 29 g／人km：採用北捷分享數值，導入時對接標準。', 'size': 12, 'sa': 2},
    {'t': '· 低門檻、低風險：以遊戲化破除 App 冷感，市府健康政策同發，實現財務自给。', 'size': 12},
])
footer(s, 11)

# ============ SLIDE 12 — 團隊 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 12, '團隊：私部門技術力 ✕ 點數模型 ✕ 公共政策')
team = [
    ('嚴世紀', 'Numbers Protocol', '專案總指揮、分散式系統架構、可信數據層驗證與信託底座', TEAL),
    ('鍾天傑', 'LINE Pay', '點數經濟模型設計、促銷乘數架構、高留存遊戲化 UX', COIN),
    ('嚴世裕', '台北市政府地政局', 'ESG 綠色政策對接、市府開放資料整合、行政合規審查與防禦', LEAF),
]
for i, (name, org, role, cl) in enumerate(team):
    x = 0.55 + i * 4.18
    c = card(s, x, 1.6, 3.85, 3.1, fill=PAPER, line_w=2.5, shadow=0.1)
    txt(s, x + 0.3, 1.9, 3.2, 0.6, [{'t': name, 'size': 22, 'bold': True}])
    chip(s, x + 0.3, 2.6, 2.4, 0.55, org, fill=cl, tcolor=(WHITE if cl != COIN else INK), size=12.5)
    txt(s, x + 0.3, 3.35, 3.25, 1.2, [{'t': role, 'size': 12.5, 'ls': 1.35}])
tech = card(s, 0.55, 5.1, 12.2, 1.55, fill=DEEP, line_w=3, shadow=0.1)
txt(s, 0.95, 5.32, 11.4, 1.15, [
    {'t': '私部門技術力 × 點數經濟實戰 × 公部門政策力', 'size': 17, 'bold': True, 'color': SUN, 'sa': 4},
    {'t': '技術支持：Omni Edge（AI 智能體平台）· Numbers Protocol（可信數據基礎建設，Human Truth. Machine Proof.）', 'size': 12, 'color': MINT},
])
footer(s, 12)

out = 'personal-auto/final-deck/從碳客變捷客_決賽簡報_v5_proposal_aligned.pptx'
prs.save(out)
print('saved', out, '| slides:', len(prs.slides.__iter__.__self__._sldIdLst))
