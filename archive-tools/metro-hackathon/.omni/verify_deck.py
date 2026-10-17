from pptx import Presentation
prs = Presentation('personal-auto/final-deck/從碳客變捷客_決賽簡報_v2.pptx')
print('slides:', len(prs.slides))
for i, slide in enumerate(prs.slides, 1):
    shapes = len(slide.shapes)
    texts = sum(1 for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip())
    pics = sum(1 for sh in slide.shapes if sh.shape_type == 13)
    first = next((sh.text_frame.text.split('\n')[0] for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()), '')
    print(f'{i:2d}: shapes={shapes:2d} texts={texts:2d} pics={pics} | {first[:40]}')
