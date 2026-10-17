# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

INK   = RGBColor(0x17,0x23,0x1C)
TEAL  = RGBColor(0x0B,0x7A,0x75)
DEEP  = RGBColor(0x08,0x3C,0x37)
LEAF  = RGBColor(0x38,0xB8,0x6F)
MINT  = RGBColor(0xC8,0xF3,0xDF)
SUN   = RGBColor(0xFF,0xD3,0x4D)
COIN  = RGBColor(0xF5,0xA5,0x24)
CORAL = RGBColor(0xFF,0x6A,0x3D)
CREAM = RGBColor(0xFF,0xF7,0xDF)
PAPER = RGBColor(0xFF,0xFD,0xF2)
GREY  = RGBColor(0x6E,0x77,0x6F)
WHITE = RGBColor(0xFF,0xFF,0xFF)
FONT  = '微軟正黑體'

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
A = 'personal-auto/final-deck/assets'
V3 = 'personal-auto/final-deck/assets/v3'

def slide_bg(s, color=CREAM):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color

def set_text(tf, lines, default_size=14, default_color=INK, align=PP_ALIGN.LEFT, bold=False, font=FONT, line_spacing=1.12):
    tf.word_wrap = True
    first = True
    for line in lines:
        if isinstance(line, str): line = {'t': line}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = line.get('align', align)
        p.line_spacing = line.get('ls', line_spacing)
        p.space_after = Pt(line.get('sa', 4))
        r = p.add_run(); r.text = line['t']
        r.font.name = line.get('font', font)
        r.font.size = Pt(line.get('size', default_size))
        r.font.bold = line.get('bold', bold)
        r.font.color.rgb = line.get('color', default_color)
        rPr = r._r.get_or_add_rPr()
        ea = rPr.find(qn('a:ea'))
        if ea is None:
            ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
        ea.set('typeface', line.get('font', font))

def card(s, x, y, w, h, fill=PAPER, line_color=INK, line_w=2.5, shadow=0.09, radius=0.12, shadow_color=INK):
    if shadow:
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x+shadow), Inches(y+shadow), Inches(w), Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = shadow_color
        sh.line.fill.background(); sh.shadow.inherit = False
        try: sh.adjustments[0] = radius
        except Exception: pass
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.fill.solid(); c.fill.fore_color.rgb = fill
    c.line.color.rgb = line_color; c.line.width = Pt(line_w)
    c.shadow.inherit = False
    try: c.adjustments[0] = radius
    except Exception: pass
    return c

def chip(s, x, y, w, h, text, fill=SUN, tcolor=INK, size=12, rot=0, bold=True, line_w=2.0):
    c = card(s, x, y, w, h, fill=fill, line_w=line_w, shadow=0.05, radius=0.5)
    if rot: c.rotation = rot
    tf = c.text_frame; tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(tf, [{'t': text, 'size': size, 'bold': bold, 'color': tcolor, 'align': PP_ALIGN.CENTER, 'sa': 0}])
    return c

def txt(s, x, y, w, h, lines, **kw):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(tb.text_frame, lines, **kw)
    return tb

def header(s, num, title, tag=None):
    bar = card(s, 0.55, 0.42, 0.62, 0.62, fill=SUN, line_w=2.5, shadow=0.07, radius=0.25)
    tfb = bar.text_frame; tfb.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(tfb, [{'t': str(num), 'size': 22, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0, 'font': 'Arial Black'}])
    txt(s, 1.4, 0.42, 9.6, 0.8, [{'t': title, 'size': 27, 'bold': True}])
    if tag:
        chip(s, 10.85, 0.5, 2.0, 0.5, tag, fill=MINT, size=12)

def footer(s, page):
    txt(s, 0.55, 7.08, 9.0, 0.35, [{'t': '從碳客變捷客 · 2026 捷運盃黑客松決賽簡報 · 想領獎金的捷客', 'size': 9, 'color': GREY}])
    txt(s, 12.2, 7.08, 0.7, 0.35, [{'t': str(page), 'size': 10, 'color': GREY, 'align': PP_ALIGN.RIGHT, 'font': 'Arial'}])

# ============ SLIDE 1 — COVER ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
card(s, 0.55, 0.55, 7.4, 6.4, fill=TEAL, line_w=3, shadow=0.12, radius=0.06)
txt(s, 1.1, 1.0, 6.4, 0.5, [{'t': '2026 捷運盃黑客松 · 決賽提案', 'size': 15, 'bold': True, 'color': MINT}])
txt(s, 1.05, 1.5, 6.6, 2.6, [
    {'t': '從碳客', 'size': 76, 'bold': True, 'color': WHITE, 'sa': 0, 'ls': 1.0},
    {'t': '變捷客', 'size': 76, 'bold': True, 'color': SUN, 'ls': 1.0},
])
hook = card(s, 1.1, 4.55, 5.9, 0.78, fill=SUN, line_w=2.5, shadow=0.08, radius=0.3)
hook.rotation = -1.5
hook.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(hook.text_frame, [{'t': '每一次通勤，都是可驗證的減碳資產', 'size': 19, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0}])
txt(s, 1.1, 5.7, 6.2, 1.0, [
    {'t': '團隊：想領獎金的捷客', 'size': 15, 'bold': True, 'color': WHITE, 'sa': 2},
    {'t': '捷運健康共享生活圈：串起市府健康政策與站內招商的永續商模', 'size': 11.5, 'color': MINT},
])
mc = card(s, 8.35, 1.15, 4.4, 3.5, fill=PAPER, line_w=3, shadow=0.12, radius=0.08)
s.shapes.add_picture(f'{A}/mascot.png', Inches(8.55), Inches(1.4), width=Inches(4.0))
chip(s, 8.6, 4.35, 1.7, 0.5, '決賽 10 強', fill=CORAL, tcolor=WHITE, size=13, rot=-2)
chip(s, 10.45, 4.35, 2.2, 0.5, '互動 Demo 可玩', fill=MINT, size=13, rot=1.5)
txt(s, 8.35, 5.35, 4.5, 1.2, [
    {'t': '線上互動原型', 'size': 12, 'bold': True, 'color': GREY, 'sa': 2},
    {'t': 'from-carbon-metro-vote.pages.dev', 'size': 13, 'bold': True, 'color': TEAL, 'font': 'Consolas'},
])
footer(s, 1)

# ============ SLIDE 2 — 痛點 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 2, '點數沒被用起來，減碳沒被感覺到')
stat = card(s, 0.55, 1.45, 12.2, 1.15, fill=DEEP, line_w=3, shadow=0.1, radius=0.1)
stat.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(stat.text_frame, [
    {'t': '臺北捷運 114 年平均日運量 207.4 萬人次 — 每一趟都是尚未被資產化的減碳貢獻', 'size': 19, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER, 'sa': 2},
    {'t': '資料來源：臺北捷運公司公開營運統計', 'size': 10.5, 'color': MINT, 'align': PP_ALIGN.CENTER, 'sa': 0},
])
pains = [
    ('Go App 資訊化', '現行 App 偏向靜態資訊查詢與單純票務服務，缺乏主動情境推播，市民日常缺乏高頻開啟的強烈動機與樂趣。', SUN),
    ('點數閒置過期', '使用者不知、亦缺乏誘因去消耗點數。捷運點大量閒置或過期逾期，導致北捷也承擔點數負債與高過期（Breakage）風險。', COIN),
    ('減碳貢獻無感', '市民知道搭乘捷運環保，但「我今天這趟省了多少碳」從未被具體量化或看見，缺乏情感連結，減碳淪為抽象概念。', MINT),
]
for i, (t, d, cl) in enumerate(pains):
    x = 0.55 + i * 4.18
    c = card(s, x, 3.0, 3.85, 2.3, fill=PAPER, line_w=2.5, shadow=0.09)
    chip(s, x + 0.25, 3.25, 1.7, 0.55, t, fill=cl, size=15)
    txt(s, x + 0.25, 4.0, 3.4, 1.2, [{'t': d, 'size': 13, 'ls': 1.25}])
bot = card(s, 0.55, 5.6, 12.2, 1.15, fill=PAPER, line_w=2.5, shadow=0.09)
bot.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(bot.text_frame, [
    {'t': '北捷視角：營運壓力大，亟需「提高 App 會員黏著度」並對接既存租賃/健康商機，活化非票務招商加值！', 'size': 15, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0},
])
footer(s, 2)

# ============ SLIDE 3 — 提案概念 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 3, '提案概念：捷運健康共享生活圈')
left_c = card(s, 0.55, 1.5, 6.0, 4.5, fill=TEAL, line_w=3, shadow=0.1)
txt(s, 0.9, 1.75, 5.3, 4.1, [
    {'t': '市政健康共享生活圈', 'size': 18, 'bold': True, 'color': SUN, 'sa': 8},
    {'t': '以「捷運點」為核心，直接對接台北市既存健康政策：', 'size': 13.5, 'bold': True, 'color': WHITE, 'sa': 8},
    {'t': '· U-Sport (體育局運動中心) 入場 → 累積減碳與健康進度', 'size': 13, 'color': MINT, 'sa': 4},
    {'t': '· U-Walk (健康局萬步挑戰) 步行 → 出站沿線完成任務', 'size': 13, 'color': MINT, 'sa': 4},
    {'t': '· YouBike (交通局接駁) 騎行 → 省碳直接加成賺點', 'size': 13, 'color': MINT, 'sa': 8},
    {'t': '★ 零新建成本：100% 串接市府既存成熟系統，免發空包彈，點數具有真實健康行為支撐！', 'size': 12.5, 'color': WHITE, 'ls': 1.25},
])
right_c = card(s, 6.85, 1.5, 5.9, 4.5, fill=PAPER, line_w=2.5, shadow=0.1)
txt(s, 7.2, 1.75, 5.2, 4.1, [
    {'t': '站內補給 ✕ 商業招商開發', 'size': 18, 'bold': True, 'color': TEAL, 'sa': 8},
    {'t': '串接站內既存租賃服務與特約廠商，轉化為高動能導客商機：', 'size': 13.5, 'bold': True, 'sa': 8},
    {'t': '· 補給：rainGo 共享傘、旅電行動電源、小樹屋電話亭會議亭', 'size': 13, 'color': GREY, 'sa': 4},
    {'t': '· 點券：亞尼克 YTM 蛋糕自助機、全家、黑潮咖啡站內櫃位', 'size': 13, 'color': GREY, 'sa': 8},
    {'t': '★ 招商加值：原本單純的租賃合約被附加「Go App 即時推播、捷運點導客券與核銷」的數位行銷方案！租賃點位單價與北捷非票務收入大幅倍增！', 'size': 13, 'ls': 1.25},
])
bar = card(s, 0.55, 6.15, 12.2, 0.78, fill=SUN, line_w=2.5, shadow=0.08)
bar.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(bar.text_frame, [{'t': '商業與市政雙贏：日常通勤賺點 ✕ 市政健康共享 ✕ 既存商業開發 $\rightarrow$ Earn-to-Spend 永續點數閉環！', 'size': 14.5, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0}])
footer(s, 3)

# ============ SLIDE 4 — 類皮克敏情感引擎 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 4, '情感引擎：類皮克敏的遊戲化世界觀', tag='內容創新')
shots = [
    (f'{V3}/story-modal-1.png', '開場故事：拯救碳森林危機', 2.55, 2.56),
    (f'{V3}/phone-hatch-modal.png', '健康任務充能 → 夥伴孵化出世', 1.90, 3.94),
    (f'{V3}/phone-forest.png', '散步小隊棲地 ✕ 分享森林卡片', 1.90, 3.94),
]
xs = [0.55, 3.45, 5.7]
for (img, cap, w, h), x in zip(shots, xs):
    ytop = 1.55 + (4.1 - h) / 2
    card(s, x, ytop - 0.12, w + 0.24, h + 0.24, fill=PAPER, line_w=2.5, shadow=0.09, radius=0.07)
    s.shapes.add_picture(img, Inches(x + 0.12), Inches(ytop), width=Inches(w))
    txt(s, x - 0.05, 5.85, w + 0.35, 0.6, [{'t': cap, 'size': 11, 'bold': True, 'align': PP_ALIGN.CENTER}])
tier_panel = card(s, 8.0, 1.55, 4.8, 4.35, fill=TEAL, line_w=3, shadow=0.1)
txt(s, 8.35, 1.78, 4.15, 4.0, [
    {'t': '以情感驅動，破除點數冷感', 'size': 17, 'bold': True, 'color': SUN, 'sa': 8},
    {'t': '將減碳數據化為生命，創造市民每天開啟的樂趣', 'size': 13, 'bold': True, 'color': WHITE, 'sa': 10, 'ls': 1.3},
    {'t': '🌲 情感世界觀 · 碳森林救援', 'size': 13, 'bold': True, 'color': MINT, 'sa': 2},
    {'t': '出站省下的碳 = 拯救濒危碳森林。小捷引導，使命感拉滿。', 'size': 12, 'color': WHITE, 'sa': 7},
    {'t': '🐣 散步夥伴與站點裝扮', 'size': 13, 'bold': True, 'color': MINT, 'sa': 2},
    {'t': 'U-Sport/Walk/Bike 任務蛋孵化夥伴「步步、輪輪、汗汗」。附帶站點限定裝扮（如市政府汗帶、忠孝復興共享傘），激發市民搜集魂。', 'size': 12, 'color': WHITE, 'sa': 7},
    {'t': '🌱 每週收成與培育神木', 'size': 13, 'bold': True, 'color': MINT, 'sa': 2},
    {'t': '週日收成儀式結算「種子」，培育神木升級。高黏著、高回訪。', 'size': 12, 'color': WHITE},
])
bar = card(s, 8.0, 6.05, 4.8, 0.75, fill=SUN, line_w=2.5, shadow=0.08)
bar.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(bar.text_frame, [{'t': '遊戲化不是點綴，是讓市民從被動配合到主動熱愛的引擎！', 'size': 13, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0}])
footer(s, 4)

# ============ SLIDE 5 — 3 分鐘 Demo 動線 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 5, '3 分鐘 Demo 動線：一條故事線走完全部機制')
journey = [
    ('① 故事與設定', '3 頁世界觀故事引入\n進行常用通勤路線 Onboarding\n（深設常用出站點）', MINT),
    ('② 出站與結算', '列車陪伴答題互動\n手動出站：+15點 / +0.23kg\n離峰通勤獲稀有種子', SUN),
    ('③ 任務與孵蛋', '做健康生活圈任務\n（YouBike/U-Sport）\n對應蛋充能 100% 孵化', COIN),
    ('④ 森林與收成', '夥伴即時入住森林棲地\n每週日儀式手動收成\n消耗種子培育神木升級', MINT),
    ('⑤ 花點換券核銷', '100點首兌全家10元券\n條碼核銷\n底層封存去識別資料指紋', SUN),
]
for i, (t, d, cl) in enumerate(journey):
    x = 0.55 + i * 2.49
    c = card(s, x, 1.6, 2.31, 3.3, fill=PAPER, line_w=2.5, shadow=0.1)
    chip(s, x + 0.18, 1.85, 1.75, 0.55, t, fill=cl, size=13)
    txt(s, x + 0.18, 2.6, 1.99, 2.1, [{'t': line, 'size': 11.5, 'ls': 1.25, 'sa': 3} for line in d.split('\n')])
b = card(s, 0.55, 5.3, 12.2, 1.5, fill=DEEP, line_w=3, shadow=0.1)
txt(s, 0.95, 5.5, 11.5, 1.1, [
    {'t': '評審亮點（第 ⑥ 步）：切到 B2B 營運主控台 — 前端乘客端的所有動作即時滾入 LIVE 互動流！', 'size': 15.5, 'bold': True, 'color': WHITE, 'sa': 3},
    {'t': '真正完成雙端聯動原型；模擬「離峰通勤」直接對映北捷「尖峰運能舒緩、離峰乘座率提升」關鍵營運 KPI。', 'size': 13, 'color': MINT},
])
footer(s, 5)

# ============ SLIDE 6 — 健康共享與站內招商 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 6, '健康共享與站內招商：無缝串接，零建成本')
sysrows = [
    ('U-Sport 體育中心', '入場運動 充能 +85% (孵化汗汗夥伴)', '市府既有健康政策'),
    ('U-Walk 萬步挑戰', '出站步行 2,000 步 充能 +65% (孵化步步)', '市府既有運動政策'),
    ('YouBike 低碳接駁', '捷運接駁 10 分鐘 充能 +35% (孵化輪輪)', '既存公共共享運具'),
    ('rainGo 共享傘 / 旅電', '站內雨天借傘 / 補電行動電源 +5 捷運點', '北捷既存租賃廠商'),
    ('亞尼克 YTM / 小樹屋', '站內蛋糕自助機 / 電話亭會議空間 享折扣', '北捷既存租賃廠商'),
]
for i, (name, val, tagx) in enumerate(sysrows):
    y = 1.55 + i * 0.92
    c = card(s, 0.55, y, 8.0, 0.78, fill=PAPER, line_w=2.5, shadow=0.07)
    txt(s, 0.9, y + 0.14, 3.4, 0.5, [{'t': name, 'size': 15, 'bold': True}])
    txt(s, 4.4, y + 0.14, 2.9, 0.5, [{'t': val, 'size': 13, 'color': TEAL, 'bold': True}])
    chip(s, 7.05, y + 0.12, 1.75, 0.52, tagx, fill=MINT, size=10.5)
side = card(s, 8.95, 1.55, 3.8, 4.55, fill=TEAL, line_w=3, shadow=0.1)
txt(s, 9.3, 1.85, 3.1, 4.0, [
    {'t': '北捷與市府的政策槓桿', 'size': 16, 'bold': True, 'color': SUN, 'sa': 8},
    {'t': '每個任務對應「已在營運」的公共設施或招商廠商。', 'size': 13, 'color': WHITE, 'sa': 6, 'ls': 1.3},
    {'t': '北捷不需新建任何硬體基礎設施，只需 API 對接與活動規則！', 'size': 13, 'color': WHITE, 'sa': 6, 'ls': 1.3},
    {'t': '租賃空間租合約被附加「Go App 推播導客能力」，單店坪效大幅提升！', 'size': 13, 'bold': True, 'color': MINT, 'sa': 8, 'ls': 1.3},
    {'t': '健康 ✕ 減碳 ✕ 商業：市府政策與北捷營收同發！', 'size': 13.5, 'bold': True, 'color': SUN},
])
footer(s, 6)

# ============ SLIDE 7 — 雙端 Demo ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 7, '雙端 LIVE 聯動：乘客端 ✕ B2B 營運主控台')
card(s, 0.7, 1.5, 3.15, 5.35, fill=PAPER, line_w=3, shadow=0.1, radius=0.07)
s.shapes.add_picture(f'{V3}/phone-today-settled.png', Inches(1.18), Inches(1.68), height=Inches(4.55))
txt(s, 0.7, 6.3, 3.15, 0.4, [{'t': '乘客端：出站結算雙入帳＋首兌進度', 'size': 11.5, 'bold': True, 'align': PP_ALIGN.CENTER}])
card(s, 4.15, 1.5, 3.15, 5.35, fill=PAPER, line_w=3, shadow=0.1, radius=0.07)
s.shapes.add_picture(f'{V3}/phone-b2b-console.png', Inches(4.63), Inches(1.68), height=Inches(4.55))
txt(s, 4.15, 6.3, 3.15, 0.4, [{'t': 'B2B 營運主控台：LIVE 互動流 ✕ 營運 KPI', 'size': 11.5, 'bold': True, 'align': PP_ALIGN.CENTER}])
info = card(s, 7.75, 1.5, 5.05, 5.35, fill=TEAL, line_w=3, shadow=0.1)
txt(s, 8.1, 1.75, 4.4, 1.7, [
    {'t': '現場一秒體驗', 'size': 20, 'bold': True, 'color': SUN, 'sa': 6},
    {'t': '純前端高擬真原型，免安裝免註冊，可在決賽現場無網完全離線執行（配備全套離線備援包）。', 'size': 13, 'color': WHITE, 'ls': 1.3},
])
qr_card = card(s, 8.1, 3.35, 2.0, 2.0, fill=PAPER, line_w=2.5, shadow=0.08, radius=0.1)
s.shapes.add_picture(f'{A}/qr-demo.png', Inches(8.22), Inches(3.47), width=Inches(1.76))
txt(s, 10.3, 3.5, 2.4, 1.9, [
    {'t': '掃碼即玩', 'size': 15, 'bold': True, 'color': WHITE, 'sa': 4},
    {'t': 'from-carbon-metro-vote.pages.dev/demo', 'size': 11.5, 'color': MINT, 'font': 'Consolas', 'ls': 1.25},
])
txt(s, 8.1, 5.75, 4.4, 0.9, [{'t': 'Demo 亮點：手按乘客端「出站、兌換」，主控台的 CPS LIVE 事件流和導客數據即時跳動，完美映射尖峰分流與招商營收！', 'size': 12, 'color': WHITE, 'ls': 1.25}])
footer(s, 7)

# ============ SLIDE 8 — Numbers 可稽核資料層 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 8, '技術硬底子：全場唯一 Numbers 可稽核數據信託')
big = card(s, 0.55, 1.5, 7.6, 3.3, fill=DEEP, line_w=3, shadow=0.1)
txt(s, 0.95, 1.75, 6.9, 2.9, [
    {'t': '發出的每一克減碳、每一點發放，', 'size': 21, 'bold': True, 'color': WHITE, 'sa': 0, 'ls': 1.15},
    {'t': '都經得起外部嚴格審計信託。', 'size': 21, 'bold': True, 'color': SUN, 'sa': 10, 'ls': 1.15},
    {'t': '· 類比雲端發票載具：市民個人減碳與核銷紀錄，自動封存為去識別化可攜資產，換手機不丟失，第三方可驗證。', 'size': 13.5, 'bold': True, 'color': WHITE, 'sa': 6, 'ls': 1.25},
    {'t': '· 資料指紋 (Provenance)：核銷與減碳憑證自動計算 SHA-256 數據指紋寫入分散式 Ledger，保證不可篡改、不重複發放。直接對接企業綠色金融或 ESG 合規審計首選！', 'size': 12, 'color': MINT, 'ls': 1.25},
])
tagc = card(s, 0.55, 5.05, 7.6, 0.85, fill=PAPER, line_w=2.5, shadow=0.08)
tagc.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(tagc.text_frame, [{'t': 'Human Truth. Machine Proof. — Numbers Protocol 支持', 'size': 16, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0, 'font': 'Consolas'}])
ph = card(s, 8.45, 1.5, 3.0, 5.1, fill=PAPER, line_w=3, shadow=0.1, radius=0.07)
s.shapes.add_picture(f'{V3}/phone-voucher-proof.png', Inches(8.9), Inches(1.68), height=Inches(4.35))
txt(s, 8.45, 6.15, 3.0, 0.45, [{'t': '核銷券附「資料指紋已封存」憑證', 'size': 11, 'bold': True, 'align': PP_ALIGN.CENTER}])
pts = card(s, 0.55, 6.1, 7.6, 0.8, fill=MINT, line_w=2.5, shadow=0.07)
pts.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(pts.text_frame, [{'t': '安全與隱私底線：Opt-in 授權、去識別化統計統計、不儲存敏感軌跡、隨時可一鍵關閉', 'size': 13, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0}])
footer(s, 8)

# ============ SLIDE 9 — 可行性與財務試算 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 9, '可行性與首年財務試算（非空中樓閣）')
nums = [('合作金流', 'NT$ 2,270 萬', TEAL), ('點數與營運支出', 'NT$ 1,525 萬', COIN), ('可支應營運餘額', 'NT$ 745 萬', LEAF)]
for i, (label, val, cl) in enumerate(nums):
    x = 0.55 + i * 4.18
    c = card(s, x, 1.55, 3.85, 1.7, fill=PAPER, line_w=2.5, shadow=0.1)
    txt(s, x + 0.3, 1.75, 3.3, 0.45, [{'t': label, 'size': 14, 'bold': True, 'color': GREY}])
    txt(s, x + 0.3, 2.2, 3.3, 0.9, [{'t': val, 'size': 28, 'bold': True, 'color': cl, 'font': 'Arial Black'}])
acct = card(s, 0.55, 3.5, 7.9, 3.3, fill=DEEP, line_w=3, shadow=0.1)
txt(s, 0.95, 3.75, 7.2, 2.9, [
    {'t': 'IFRS 15 點數會計與 Breakage 管理（直球對決財務評審）', 'size': 16, 'bold': True, 'color': SUN, 'sa': 8},
    {'t': '· 發點遞延處理：發放時列合約負債（遞延收入），實際核銷或逾期時認列收入。', 'size': 13, 'color': WHITE, 'sa': 5, 'ls': 1.2},
    {'t': '· Breakage 精精算：引入 OPENPOINT 歷史未兌換率（Breakage）估算，不虛增負債壓力。', 'size': 13, 'color': WHITE, 'sa': 5, 'ls': 1.2},
    {'t': '· 成本共同分攤：站內特約商家、既存租賃商共同出資點數池，北捷負債壓力完全可控！', 'size': 13, 'color': WHITE, 'sa': 5, 'ls': 1.2},
    {'t': '· 研究底稿：完整分析 LINE POINTS、星巴克與 Marriott 之會計實務。', 'size': 11.5, 'color': MINT},
])
src = card(s, 8.75, 3.5, 4.0, 3.3, fill=PAPER, line_w=2.5, shadow=0.1)
txt(s, 9.05, 3.75, 3.5, 2.9, [
    {'t': '金流來源拆解', 'size': 15, 'bold': True, 'color': TEAL, 'sa': 6},
    {'t': '· 商家精準廣告與導客分成 (5%–15%)', 'size': 12.5, 'sa': 4},
    {'t': '· 站內租賃空間 App 導客溢價租金', 'size': 12.5, 'sa': 4},
    {'t': '· 市府健康/綠色生活圈補貼', 'size': 12.5, 'sa': 4},
    {'t': '· 綠色企業 ESG 數據稽核與採購分成', 'size': 12.5},
])
footer(s, 9)

# ============ SLIDE 10 — 藍圖與誠實標註 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 10, '三階段導入藍圖與誠實標註')
phases = [
    ('決賽階段 (現有成果)', '可操作高擬真原型與 B2B 主控台，100% 離線備援可用，展現核心商業與健康機制。', MINT),
    ('試辦階段 (半年導入)', '選定 2–3 個示範生活圈站點，對接 U-Sport / Walk 政策，串聯站內既存廠商出站導客。', SUN),
    ('擴大階段 (一年推廣)', '活化常客優惠轉點機制，對接悠遊付百萬商家網絡，引進綠色企業 ESG 數據信託採購。', COIN),
]
for i, (t, d, cl) in enumerate(phases):
    x = 0.55 + i * 4.18
    c = card(s, x, 1.55, 3.85, 2.5, fill=PAPER, line_w=2.5, shadow=0.1)
    chip(s, x + 0.25, 1.75, 2.1, 0.52, t, fill=cl, size=13.5)
    txt(s, x + 0.28, 2.45, 3.3, 1.5, [{'t': d, 'size': 12.5, 'ls': 1.25}])
    if i < 2:
        txt(s, x + 3.85, 2.4, 0.4, 0.8, [{'t': '→', 'size': 18, 'bold': True, 'align': PP_ALIGN.CENTER}])
real = card(s, 0.55, 4.35, 6.0, 2.5, fill=PAPER, line_w=2.5, shadow=0.08)
txt(s, 0.9, 4.55, 5.3, 2.1, [
    {'t': '✓ 100% 驗證 / 有依據的真實資產', 'size': 14.5, 'bold': True, 'color': LEAF, 'sa': 4},
    {'t': '· 北捷 207.4 萬日運量、雨天 rainGo 共享傘、YouBike 既存接駁。', 'size': 12, 'sa': 2},
    {'t': '· U-Sport / Walk 台北市體育局與衛生局既有成熟健康系統。', 'size': 12, 'sa': 2},
    {'t': '· 採用臺北市資料大平臺「各站分時進出量」開放資料進行時段設計。', 'size': 12},
])
todo = card(s, 6.75, 4.35, 6.0, 2.5, fill=PAPER, line_w=2.5, shadow=0.08)
txt(s, 7.1, 4.55, 5.3, 2.1, [
    {'t': '△ 待北捷核准 / 落地調整之假設（誠實標註）', 'size': 14.5, 'bold': True, 'color': CORAL, 'sa': 4},
    {'t': '· 減碳係數 104 / 75 / 29 g／人km（北捷座談分享，導入時對接標準）。', 'size': 12, 'sa': 2},
    {'t': '· 財務精算中所涉及之月活會員（Go App 會員基底約 20-30 萬）與導客率。', 'size': 12, 'sa': 2},
    {'t': '· 試辦階段之具體分流加碼點數、團隊與北捷之技術授權與營運義務。', 'size': 12},
])
footer(s, 10)

# ============ SLIDE 11 — 團隊 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 11, '團隊：私部門技術力 ✕ 點數模型 ✕ 公共政策')
team = [
    ('嚴世紀', 'Numbers Protocol', '專案總指揮、分散式系統架構、可信數據層驗證與信託底座', TEAL),
    ('鍾天傑', 'LINE Pay', '點數經濟模型設計、促銷乘數架構、高留存遊戲化 UX', COIN),
    ('嚴世裕', '台北市政府地政局', 'ESG 綠色政策對接、市府開放資料整合、行政合規審查與防禦', LEAF),
]
for i, (name, org, role, cl) in enumerate(team):
    x = 0.55 + i * 4.18
    c = card(s, x, 1.6, 3.85, 3.1, fill=PAPER, line_w=2.5, shadow=0.1)
    txt(s, x + 0.3, 1.9, 3.2, 0.6, [{'t': name, 'size': 22, 'bold': True}])
    chip(s, x + 0.3, 2.6, 2.4, 0.55, org, fill=cl, tcolor=(WHITE if cl != COIN else INK), size=12.5)
    txt(s, x + 0.3, 3.35, 3.25, 1.2, [{'t': role, 'size': 12.5, 'ls': 1.35}])
tech = card(s, 0.55, 5.1, 12.2, 1.55, fill=DEEP, line_w=3, shadow=0.1)
txt(s, 0.95, 5.32, 11.4, 1.15, [
    {'t': '私部門技術力 × 點數經濟實戰 × 公部門政策力', 'size': 17, 'bold': True, 'color': SUN, 'sa': 4},
    {'t': '技術支持：Omni Edge（AI 智能體平台）· Numbers Protocol（可信數據基礎建設，Human Truth. Machine Proof.）', 'size': 12, 'color': MINT},
])
footer(s, 11)

# ============ SLIDE 12 — 附錄 QA 與收尾 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 12, '附錄 QA 備彈與收尾')
left_box = card(s, 0.55, 1.5, 7.3, 4.8, fill=PAPER, line_w=2.5, shadow=0.1)
txt(s, 0.85, 1.7, 6.7, 4.4, [
    {'t': '評審 QA 直球回答（黑傑客備彈）', 'size': 16, 'bold': True, 'color': TEAL, 'sa': 10},
    {'t': 'Q：點數發放與回饋成本誰出？會不會拖累北捷財務？', 'size': 12.5, 'bold': True, 'color': CORAL, 'sa': 2},
    {'t': 'A：由站內既存租賃商家與全家等特約商之「行銷與廣告預算」出資，北捷非票務招商營收增加而非成本增加，財務完全可控。', 'size': 11.5, 'sa': 6, 'ls': 1.15},
    {'t': 'Q：被挑戰減碳數字不準或灌水怎麼辦？', 'size': 12.5, 'bold': True, 'color': CORAL, 'sa': 2},
    {'t': 'A：29 gCO2e/km 為可配置參數（Demo 可調），採用北捷座談會數值並誠實標註。導入時100%對接環境部或北捷標準。數據均附 Numbers 資料指紋防篡改。', 'size': 11.5, 'sa': 6, 'ls': 1.15},
    {'t': 'Q：跟 Go App 整合需要改動核心系統嗎？成本多高？', 'size': 12.5, 'bold': True, 'color': CORAL, 'sa': 2},
    {'t': 'A：不改動售票/進出站核心，以「輕量 API 對接推播與 Go!Map 商家庫」做加值模組。試辦僅需 2–3 站既存生活圈與現有廠商，導入門檻與風險極低。', 'size': 11.5, 'ls': 1.15},
])
right_box = card(s, 8.1, 1.5, 4.65, 4.8, fill=TEAL, line_w=3, shadow=0.1)
txt(s, 8.4, 1.75, 4.1, 4.3, [
    {'t': '每一次通勤，都是可驗證的減碳資產', 'size': 17, 'bold': True, 'color': SUN, 'sa': 6},
    {'t': '從碳客變捷客 — 讓每趟捷運，成為省錢與減碳的情感證明', 'size': 12.5, 'color': MINT, 'sa': 15, 'ls': 1.2},
    {'t': '互動原型（免安裝 · 可完全離線體驗）', 'size': 12.5, 'bold': True, 'color': WHITE, 'sa': 4},
    {'t': 'https://from-carbon-metro-vote.pages.dev/demo/', 'size': 11, 'bold': True, 'color': MINT, 'font': 'Consolas', 'sa': 15},
    {'t': '大會人氣票選進行中，投給【從碳客變捷客】！', 'size': 12, 'color': WHITE},
])
fin = card(s, 8.35, 5.0, 4.15, 1.1, fill=SUN, line_w=2.5, shadow=0.08)
fin.rotation = -1.5
fin.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(fin.text_frame, [
    {'t': '黑捷客出發，創意進站！', 'size': 14.5, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 2},
    {'t': '我們決賽現場見！🏆', 'size': 14.5, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0},
])
footer(s, 12)

out = 'personal-auto/final-deck/從碳客變捷客_決賽簡報_v3.pptx'
prs.save(out)
print('saved', out, '| slides:', len(prs.slides.__iter__.__self__._sldIdLst))
