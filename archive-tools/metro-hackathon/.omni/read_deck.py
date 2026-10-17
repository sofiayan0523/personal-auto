from pptx import Presentation
prs = Presentation('personal-auto/final-deck/decksrc_v1.pptx')
print('slides:', len(prs.slides), 'size:', prs.slide_width, prs.slide_height)
for i, slide in enumerate(prs.slides, 1):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip().replace('\n', ' / ')
            if t: texts.append(t)
    print(f'--- {i} ---')
    print(' | '.join(texts)[:600])
