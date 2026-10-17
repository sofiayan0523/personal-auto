# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

INK   = RGBColor(0x17,0x23,0x1C)
TEAL  = RGBColor(0x0B,0x7A,0x75)
DEEP  = RGBColor(0x08,0x3C,0x37)
LEAF  = RGBColor(0x38,0xB8,0x6F)
MINT  = RGBColor(0xC8,0xF3,0xDF)
SUN   = RGBColor(0xFF,0xD3,0x4D)
COIN  = RGBColor(0xF5,0xA5,0x24)
CORAL = RGBColor(0xFF,0x6A,0x3D)
CREAM = RGBColor(0xFF,0xF7,0xDF)
PAPER = RGBColor(0xFF,0xFD,0xF2)
GREY  = RGBColor(0x6E,0x77,0x6F)
WHITE = RGBColor(0xFF,0xFF,0xFF)
FONT  = '微軟正黑體'

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
A = 'personal-auto/final-deck/assets'

def slide_bg(s, color=CREAM):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color

def set_text(tf, lines, default_size=14, default_color=INK, align=PP_ALIGN.LEFT, bold=False, font=FONT, line_spacing=1.12):
    tf.word_wrap = True
    first = True
    for line in lines:
        if isinstance(line, str): line = {'t': line}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = line.get('align', align)
        p.line_spacing = line.get('ls', line_spacing)
        p.space_after = Pt(line.get('sa', 4))
        r = p.add_run(); r.text = line['t']
        r.font.name = line.get('font', font)
        r.font.size = Pt(line.get('size', default_size))
        r.font.bold = line.get('bold', bold)
        r.font.color.rgb = line.get('color', default_color)
        ea = r._r.rPr.get_or_change_to_latin() if False else None
        # set East Asian font
        rPr = r._r.get_or_add_rPr()
        ea = rPr.find(qn('a:ea'))
        if ea is None:
            ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
        ea.set('typeface', line.get('font', font))

def card(s, x, y, w, h, fill=PAPER, line_color=INK, line_w=2.5, shadow=0.09, radius=0.12, shadow_color=INK):
    if shadow:
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x+shadow), Inches(y+shadow), Inches(w), Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = shadow_color
        sh.line.fill.background(); sh.shadow.inherit = False
        try: sh.adjustments[0] = radius
        except Exception: pass
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.fill.solid(); c.fill.fore_color.rgb = fill
    c.line.color.rgb = line_color; c.line.width = Pt(line_w)
    c.shadow.inherit = False
    try: c.adjustments[0] = radius
    except Exception: pass
    return c

def chip(s, x, y, w, h, text, fill=SUN, tcolor=INK, size=12, rot=0, bold=True, line_w=2.0):
    c = card(s, x, y, w, h, fill=fill, line_w=line_w, shadow=0.05, radius=0.5)
    if rot: c.rotation = rot
    tf = c.text_frame; tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(tf, [{'t': text, 'size': size, 'bold': bold, 'color': tcolor, 'align': PP_ALIGN.CENTER, 'sa': 0}])
    return c

def txt(s, x, y, w, h, lines, **kw):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(tb.text_frame, lines, **kw)
    return tb

def header(s, num, title, tag=None):
    bar = card(s, 0.55, 0.42, 0.62, 0.62, fill=SUN, line_w=2.5, shadow=0.07, radius=0.25)
    tfb = bar.text_frame; tfb.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(tfb, [{'t': str(num), 'size': 22, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0, 'font': 'Arial Black'}])
    txt(s, 1.4, 0.42, 9.6, 0.8, [{'t': title, 'size': 27, 'bold': True}])
    if tag:
        chip(s, 10.85, 0.5, 2.0, 0.5, tag, fill=MINT, size=12)

def footer(s, page):
    txt(s, 0.55, 7.08, 9.0, 0.35, [{'t': '從碳客變捷客 · 2026 捷運盃黑客松決賽 · 想領獎金的捷客', 'size': 9, 'color': GREY}])
    txt(s, 12.2, 7.08, 0.7, 0.35, [{'t': str(page), 'size': 10, 'color': GREY, 'align': PP_ALIGN.RIGHT, 'font': 'Arial'}])

# ============ SLIDE 1 — COVER ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
# big teal panel
card(s, 0.55, 0.55, 7.4, 6.4, fill=TEAL, line_w=3, shadow=0.12, radius=0.06)
txt(s, 1.1, 1.0, 6.4, 0.5, [{'t': '2026 捷運盃黑客松 · 決賽', 'size': 15, 'bold': True, 'color': MINT}])
txt(s, 1.05, 1.5, 6.6, 2.6, [
    {'t': '從碳客', 'size': 76, 'bold': True, 'color': WHITE, 'sa': 0, 'ls': 1.0},
    {'t': '變捷客', 'size': 76, 'bold': True, 'color': SUN, 'ls': 1.0},
])
hook = card(s, 1.1, 4.55, 5.9, 0.78, fill=SUN, line_w=2.5, shadow=0.08, radius=0.3)
hook.rotation = -1.5
hook.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(hook.text_frame, [{'t': '每一次通勤，都是可驗證的減碳資產', 'size': 19, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0}])
txt(s, 1.1, 5.7, 6.2, 1.0, [
    {'t': '團隊：想領獎金的捷客', 'size': 15, 'bold': True, 'color': WHITE, 'sa': 2},
    {'t': '整合於台北捷運 Go App 的 AI 智慧服務模組', 'size': 12.5, 'color': MINT},
])
# mascot card
mc = card(s, 8.35, 1.15, 4.4, 3.5, fill=PAPER, line_w=3, shadow=0.12, radius=0.08)
s.shapes.add_picture(f'{A}/mascot.png', Inches(8.55), Inches(1.4), width=Inches(4.0))
chip(s, 8.6, 4.35, 1.7, 0.5, '決賽 10 強', fill=CORAL, tcolor=WHITE, size=13, rot=-2)
chip(s, 10.45, 4.35, 2.2, 0.5, '互動 Demo 可玩', fill=MINT, size=13, rot=1.5)
txt(s, 8.35, 5.35, 4.5, 1.2, [
    {'t': '線上互動原型', 'size': 12, 'bold': True, 'color': GREY, 'sa': 2},
    {'t': 'from-carbon-metro-vote.pages.dev', 'size': 13, 'bold': True, 'color': TEAL, 'font': 'Consolas'},
])

# ============ SLIDE 2 — 痛點 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 2, '點數沒被用起來，減碳沒被感覺到')
stat = card(s, 0.55, 1.45, 12.2, 1.15, fill=DEEP, line_w=3, shadow=0.1, radius=0.1)
stat.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(stat.text_frame, [
    {'t': '臺北捷運 114 年平均日運量 207.4 萬人次 — 每一趟都是尚未被資產化的減碳貢獻', 'size': 19, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER, 'sa': 2},
    {'t': '資料來源：臺北捷運公司公開營運統計', 'size': 10.5, 'color': MINT, 'align': PP_ALIGN.CENTER, 'sa': 0},
])
pains = [
    ('點數閒置', '使用者須自行開地圖找店、缺乏主動推薦，大量捷運點閒置或逾期', SUN),
    ('推播無感', 'Go App 資訊量大、banner 易被滑過，個人化程度不足', COIN),
    ('減碳無感', '大家知道搭捷運環保，但「我今天減了多少碳」從未被看見', MINT),
]
for i, (t, d, cl) in enumerate(pains):
    x = 0.55 + i * 4.18
    c = card(s, x, 3.0, 3.85, 2.3, fill=PAPER, line_w=2.5, shadow=0.09)
    chip(s, x + 0.25, 3.25, 1.6, 0.55, t, fill=cl, size=15)
    txt(s, x + 0.25, 4.0, 3.4, 1.2, [{'t': d, 'size': 13.5, 'ls': 1.25}])
bot = card(s, 0.55, 5.6, 12.2, 1.15, fill=PAPER, line_w=2.5, shadow=0.09)
bot.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(bot.text_frame, [
    {'t': '北捷視角：本業營運壓力大，需要「能變現的會員經營」與「非票務收入」工具', 'size': 16, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0},
])
footer(s, 2)

# ============ SLIDE 3 — 競爭定位 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 3, '10 強都在發點數，只有我們做可驗證的減碳資產', tag='競爭定位')
rows = [
    ('點數賺花閉環', '7 / 8 隊', '✓ 有，且與減碳雙向聯動'),
    ('店家導客 O2O', '7 / 8 隊', '✓ 有，串北捷既有招商業務'),
    ('遊戲化任務', '6 / 8 隊', '✓ 有，乘數 Level Up'),
    ('減碳量化為核心機制', '0 隊', '★ 獨家：104/75/29 係數 × 點數乘數'),
    ('資料可稽核（provenance）', '0 隊', '★ 獨家：Numbers Protocol 資料指紋'),
    ('IFRS 15 點數負債論述', '0 隊', '★ 獨家：遞延收入 + breakage 估算'),
]
card(s, 0.55, 1.5, 12.2, 4.55, fill=PAPER, line_w=2.5, shadow=0.1)
txt(s, 0.95, 1.7, 5.2, 0.4, [{'t': '構面', 'size': 13, 'bold': True, 'color': GREY}])
txt(s, 6.2, 1.7, 2.3, 0.4, [{'t': '其他隊伍', 'size': 13, 'bold': True, 'color': GREY}])
txt(s, 8.55, 1.7, 4.0, 0.4, [{'t': '從碳客變捷客', 'size': 13, 'bold': True, 'color': TEAL}])
for i, (dim, other, ours) in enumerate(rows):
    y = 2.15 + i * 0.63
    star = ours.startswith('★')
    txt(s, 0.95, y, 5.1, 0.55, [{'t': dim, 'size': 14.5, 'bold': star}])
    txt(s, 6.2, y, 2.3, 0.55, [{'t': other, 'size': 13.5, 'color': GREY}])
    txt(s, 8.55, y, 4.0, 0.55, [{'t': ours, 'size': 13.5, 'bold': star, 'color': (CORAL if star else INK)}])
concl = card(s, 0.55, 6.25, 12.2, 0.72, fill=SUN, line_w=2.5, shadow=0.08)
concl.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(concl.text_frame, [{'t': '紅海不打：我們的戰場是「數據可信 × 全串既有系統 × 財務可落地」', 'size': 16, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0}])
footer(s, 3)

# ============ SLIDE 4 — 核心機制 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 4, '核心機制：捷運點 × 減碳量 雙向聯動')
f = card(s, 0.55, 1.5, 6.0, 3.1, fill=TEAL, line_w=3, shadow=0.1)
txt(s, 0.9, 1.72, 5.4, 0.45, [{'t': '碳節省公式（係數可配置）', 'size': 14, 'bold': True, 'color': MINT}])
txt(s, 0.9, 2.2, 5.4, 1.3, [
    {'t': '(104 − 75) × 里程 = 每公里省 29 g', 'size': 24, 'bold': True, 'color': WHITE, 'sa': 4},
    {'t': '小客車 104 g／人km － 捷運 75 g／人km', 'size': 13, 'color': MINT},
])
txt(s, 0.9, 3.75, 5.4, 0.8, [
    {'t': '· 北捷提供之參考值，正式導入前由北捷確認', 'size': 11.5, 'color': MINT, 'sa': 2},
    {'t': '· 不另立「綠點」：加成直接作用於捷運點', 'size': 11.5, 'color': MINT},
])
m = card(s, 6.85, 1.5, 5.9, 3.1, fill=PAPER, line_w=2.5, shadow=0.1)
txt(s, 7.2, 1.72, 5.2, 0.45, [{'t': '月度減碳乘數 Level Up', 'size': 14, 'bold': True, 'color': TEAL}])
tiers = [('0–5 kg', '1.0×', MINT), ('5–10 kg', '1.2×', SUN), ('10 kg+', '1.5×', COIN)]
for i, (rng, mult, cl) in enumerate(tiers):
    x = 7.2 + i * 1.85
    chip(s, x, 2.35, 1.6, 1.0, mult, fill=cl, size=24, line_w=2.5)
    txt(s, x, 3.5, 1.6, 0.4, [{'t': rng, 'size': 12, 'align': PP_ALIGN.CENTER, 'color': GREY}])
txt(s, 7.2, 4.0, 5.2, 0.5, [{'t': '減碳越多 → 賺點乘數越高 → 更想搭捷運', 'size': 13.5, 'bold': True}])
loop = card(s, 0.55, 4.95, 12.2, 1.85, fill=PAPER, line_w=2.5, shadow=0.1)
txt(s, 0.9, 5.15, 11.5, 0.45, [{'t': '完整閉環（Demo 可實際操作）', 'size': 14, 'bold': True, 'color': TEAL}])
steps = ['搭捷運 賺點+減碳', '情境任務 乘數加成', '出站推播 導客', '花點換券 核銷', '減碳證明 可稽核']
for i, st in enumerate(steps):
    x = 0.9 + i * 2.42
    chip(s, x, 5.75, 2.1, 0.75, st, fill=(MINT if i % 2 == 0 else SUN), size=12.5)
    if i < 4:
        txt(s, x + 2.08, 5.75, 0.4, 0.75, [{'t': '→', 'size': 16, 'bold': True, 'align': PP_ALIGN.CENTER}])
footer(s, 4)

# ============ SLIDE 5 — 使用者旅程 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 5, '使用者旅程（決賽 Demo 動線）')
journey = [
    ('1 進站前', '3 題 opt-in 偏好\nonboarding，不假設\n北捷已掌握消費偏好', MINT),
    ('2 通勤中', '完成通勤：+18 點\n+0.23 kg 減碳\n雙數字同時入帳', SUN),
    ('3 出站時', '出站推播：3 分鐘內\n2 家可用捷運點店家\n雨天自動推借傘任務', COIN),
    ('4 生活圈', 'App 換券 → 條碼核銷\n借傘/補電/YTM/運動\n每筆附資料指紋', MINT),
]
for i, (t, d, cl) in enumerate(journey):
    x = 0.55 + i * 3.13
    c = card(s, x, 1.6, 2.85, 3.3, fill=PAPER, line_w=2.5, shadow=0.1)
    chip(s, x + 0.22, 1.85, 1.5, 0.55, t, fill=cl, size=14)
    txt(s, x + 0.22, 2.6, 2.45, 2.1, [{'t': line, 'size': 12.5, 'ls': 1.25, 'sa': 3} for line in d.split('\n')])
b = card(s, 0.55, 5.3, 12.2, 1.5, fill=DEEP, line_w=3, shadow=0.1)
txt(s, 0.95, 5.5, 11.5, 1.1, [
    {'t': '第 5 步（評審亮點）：切到北捷 Dashboard — 剛才每個動作即時滾入 LIVE 事件流', 'size': 16, 'bold': True, 'color': WHITE, 'sa': 3},
    {'t': '乘客端與營運端同一套資料聯動，閉環親眼可見', 'size': 13, 'color': MINT},
])
footer(s, 5)

# ============ SLIDE 6 — 健康共享生活圈 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 6, '健康共享生活圈：串市府既有系統，零新建')
sysrows = [
    ('U-Sport 運動中心', '入場運動 +25 點', '臺北市既有系統'),
    ('U-Walk 萬步挑戰', '出站補 2,000 步 +20 點', '臺北市既有系統'),
    ('YouBike 低碳接駁', '捷運後 1–2 km 接駁 +12 點', '既有共享運具'),
    ('raingo 共享傘', '雨天出站借傘 +5 點', '北捷站內既有廠商'),
    ('旅電 ChargeSPOT', '共享行動電源 +5 點', '北捷站內既有廠商'),
]
for i, (name, val, tagx) in enumerate(sysrows):
    y = 1.55 + i * 0.92
    c = card(s, 0.55, y, 8.0, 0.78, fill=PAPER, line_w=2.5, shadow=0.07)
    txt(s, 0.9, y + 0.14, 3.4, 0.5, [{'t': name, 'size': 15, 'bold': True}])
    txt(s, 4.4, y + 0.14, 2.9, 0.5, [{'t': val, 'size': 13, 'color': TEAL, 'bold': True}])
    chip(s, 7.05, y + 0.12, 1.75, 0.52, tagx, fill=MINT, size=10.5)
side = card(s, 8.95, 1.55, 3.8, 4.55, fill=TEAL, line_w=3, shadow=0.1)
txt(s, 9.3, 1.85, 3.1, 4.0, [
    {'t': '為什麼可行性滿分', 'size': 16, 'bold': True, 'color': SUN, 'sa': 8},
    {'t': '每個任務都對應「已在營運」的系統或廠商', 'size': 13, 'color': WHITE, 'sa': 6, 'ls': 1.3},
    {'t': '北捷不需新建任何基礎設施，只需 API 對接與活動規則', 'size': 13, 'color': WHITE, 'sa': 6, 'ls': 1.3},
    {'t': '健康 × 減碳 × 點數：市府政策三箭同發', 'size': 13, 'bold': True, 'color': MINT, 'ls': 1.3},
])
footer(s, 6)

# ============ SLIDE 7 — 站內招商加值 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 7, '站內招商加值：把場站商業開發變成新營收線')
txt(s, 0.55, 1.45, 12.2, 0.5, [{'t': '北捷站內既存租賃點位（皆為公開可查的真實案例），全部可被 Go App 推播、Go!Map 標記、捷運點券導客：', 'size': 14}])
vendors = ['亞尼克 YTM 蛋糕機', 'raingo 共享傘', '旅電 / ChargeSPOT', 'Happ. 小樹屋電話亭', 'TRICENT 試香機', 'Yo-Kai 拉麵機', '客製悠遊卡機', '快剪 / 快閃櫃']
for i, v in enumerate(vendors):
    x = 0.55 + (i % 4) * 3.13
    y = 2.15 + (i // 4) * 0.85
    chip(s, x, y, 2.85, 0.65, v, fill=(SUN if i % 2 == 0 else MINT), size=13)
val = card(s, 0.55, 4.1, 12.2, 2.55, fill=PAPER, line_w=2.5, shadow=0.1)
txt(s, 0.95, 4.32, 11.5, 2.1, [
    {'t': '對北捷的三重價值', 'size': 15, 'bold': True, 'color': TEAL, 'sa': 6},
    {'t': '1) 招商加值：租賃點位附帶「App 導客能力」出租，點位單價與續約率提升', 'size': 14, 'sa': 4},
    {'t': '2) 精準變現：某站 × 某時段 × 某偏好的推播，廣告與行銷方案賣的是成效', 'size': 14, 'sa': 4},
    {'t': '3) 邊界清楚：租賃廠商自有行銷與第三方廣告仍依北捷既有審查規範', 'size': 14},
])
footer(s, 7)

# ============ SLIDE 8 — 雙端 Demo ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 8, '雙端 LIVE Demo：乘客端 × 北捷 Dashboard')
card(s, 0.7, 1.5, 3.15, 5.35, fill=PAPER, line_w=3, shadow=0.1, radius=0.07)
s.shapes.add_picture(f'{A}/phone-demo.png', Inches(0.86), Inches(1.68), height=Inches(4.55))
txt(s, 0.7, 6.3, 3.15, 0.4, [{'t': '乘客端：賺點＋減碳雙入帳', 'size': 11.5, 'bold': True, 'align': PP_ALIGN.CENTER}])
card(s, 4.15, 1.5, 3.15, 5.35, fill=PAPER, line_w=3, shadow=0.1, radius=0.07)
s.shapes.add_picture(f'{A}/phone-dashboard.png', Inches(4.31), Inches(1.68), height=Inches(4.55))
txt(s, 4.15, 6.3, 3.15, 0.4, [{'t': 'Dashboard：LIVE 事件流即時聯動', 'size': 11.5, 'bold': True, 'align': PP_ALIGN.CENTER}])
info = card(s, 7.75, 1.5, 5.05, 5.35, fill=TEAL, line_w=3, shadow=0.1)
txt(s, 8.1, 1.75, 4.4, 1.7, [
    {'t': '現場就能玩', 'size': 20, 'bold': True, 'color': SUN, 'sa': 6},
    {'t': '純前端原型・免安裝免註冊・可完全離線執行（附離線備援包）', 'size': 13, 'color': WHITE, 'ls': 1.3},
])
qr_card = card(s, 8.1, 3.35, 2.0, 2.0, fill=PAPER, line_w=2.5, shadow=0.08, radius=0.1)
s.shapes.add_picture(f'{A}/qr-demo.png', Inches(8.22), Inches(3.47), width=Inches(1.76))
txt(s, 10.3, 3.5, 2.4, 1.9, [
    {'t': '掃碼體驗', 'size': 15, 'bold': True, 'color': WHITE, 'sa': 4},
    {'t': 'from-carbon-metro-vote.pages.dev/demo', 'size': 11.5, 'color': MINT, 'font': 'Consolas', 'ls': 1.25},
])
txt(s, 8.1, 5.75, 4.4, 0.9, [{'t': 'Demo 亮點：完成借傘任務 → 切 Dashboard → 觸及/導客/減碳數字即時跳動', 'size': 12.5, 'color': WHITE, 'ls': 1.3}])
footer(s, 8)

# ============ SLIDE 9 — Numbers 可稽核資料層 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 9, '全場唯一：Numbers 可稽核資料層')
big = card(s, 0.55, 1.5, 7.6, 3.3, fill=DEEP, line_w=3, shadow=0.1)
txt(s, 0.95, 1.75, 6.9, 2.9, [
    {'t': '北捷發出去的每一點、', 'size': 24, 'bold': True, 'color': WHITE, 'sa': 0, 'ls': 1.15},
    {'t': '每一克減碳，都經得起審計。', 'size': 24, 'bold': True, 'color': SUN, 'sa': 10, 'ls': 1.15},
    {'t': '減碳統計與核銷紀錄以去識別化方式產生資料指紋（provenance），寫入不可竄改的可驗證紀錄 — ESG 揭露與金融綠色回饋可直接引用', 'size': 13.5, 'color': MINT, 'ls': 1.35},
])
tagc = card(s, 0.55, 5.05, 7.6, 0.85, fill=PAPER, line_w=2.5, shadow=0.08)
tagc.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(tagc.text_frame, [{'t': 'Human Truth. Machine Proof. — Numbers Protocol', 'size': 16, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0, 'font': 'Consolas'}])
ph = card(s, 8.45, 1.5, 3.0, 5.1, fill=PAPER, line_w=3, shadow=0.1, radius=0.07)
s.shapes.add_picture(f'{A}/phone-voucher.png', Inches(8.6), Inches(1.68), height=Inches(4.35))
txt(s, 8.45, 6.15, 3.0, 0.45, [{'t': '核銷券附「資料指紋已封存」', 'size': 11, 'bold': True, 'align': PP_ALIGN.CENTER}])
pts = card(s, 0.55, 6.1, 7.6, 0.8, fill=MINT, line_w=2.5, shadow=0.07)
pts.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(pts.text_frame, [{'t': '隱私底線：opt-in、僅去識別化月度統計值、不儲存起訖站與身分、可隨時關閉', 'size': 13, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0}])
footer(s, 9)

# ============ SLIDE 10 — 可行性對照表 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 10, '可行性對照：不是另一支 App，是 Go App 的加值模組', tag='可行性')
feat = [
    ('碳計算與乘數', '票務進出站紀錄 + 官方係數', '純演算法，零新硬體'),
    ('出站推播導客', 'Go App 推播 + Go!Map 商家資料', '輕量 API，不動核心'),
    ('花點換券核銷', '既有捷運點 + App 券 + 掃碼', '不假設 POS 全串接'),
    ('健康共享任務', 'U-Sport / U-Walk / YouBike', '市府既有系統對接'),
    ('站內補給任務', 'raingo / 旅電 / YTM / 小樹屋', '既存租賃廠商'),
    ('可稽核資料層', 'Numbers Protocol API', '僅上傳去識別統計值'),
]
card(s, 0.55, 1.5, 12.2, 4.6, fill=PAPER, line_w=2.5, shadow=0.1)
txt(s, 0.95, 1.68, 3.6, 0.4, [{'t': '功能', 'size': 13, 'bold': True, 'color': GREY}])
txt(s, 4.6, 1.68, 4.4, 0.4, [{'t': '依賴的既有系統', 'size': 13, 'bold': True, 'color': GREY}])
txt(s, 9.1, 1.68, 3.4, 0.4, [{'t': '導入需求', 'size': 13, 'bold': True, 'color': GREY}])
for i, (a, b, c_) in enumerate(feat):
    y = 2.12 + i * 0.66
    txt(s, 0.95, y, 3.6, 0.55, [{'t': a, 'size': 14, 'bold': True}])
    txt(s, 4.6, y, 4.5, 0.55, [{'t': b, 'size': 13}])
    txt(s, 9.1, y, 3.5, 0.55, [{'t': c_, 'size': 13, 'color': TEAL, 'bold': True}])
con = card(s, 0.55, 6.25, 12.2, 0.72, fill=SUN, line_w=2.5, shadow=0.08)
con.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(con.text_frame, [{'t': '每一項功能都有「已在營運」的依賴系統 — 北捷核准後即可開工', 'size': 16, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0}])
footer(s, 10)

# ============ SLIDE 11 — 財務試算 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 11, '首年財務試算 + 點數會計（全場唯一）', tag='待北捷確認')
nums = [('合作金流', 'NT$ 2,270 萬', TEAL), ('點數與營運支出', 'NT$ 1,525 萬', COIN), ('可支應營運餘額', 'NT$ 745 萬', LEAF)]
for i, (label, val, cl) in enumerate(nums):
    x = 0.55 + i * 4.18
    c = card(s, x, 1.55, 3.85, 1.9, fill=PAPER, line_w=2.5, shadow=0.1)
    txt(s, x + 0.3, 1.8, 3.3, 0.45, [{'t': label, 'size': 14, 'bold': True, 'color': GREY}])
    txt(s, x + 0.3, 2.3, 3.3, 0.9, [{'t': val, 'size': 30, 'bold': True, 'color': cl, 'font': 'Arial Black'}])
acct = card(s, 0.55, 3.85, 7.9, 2.8, fill=DEEP, line_w=3, shadow=0.1)
txt(s, 0.95, 4.1, 7.2, 2.4, [
    {'t': 'IFRS 15 點數負債處理（評審 QA 直球回答）', 'size': 16, 'bold': True, 'color': SUN, 'sa': 8},
    {'t': '· 發點時列遞延收入（合約負債），核銷或到期才認列收入', 'size': 13.5, 'color': WHITE, 'sa': 5},
    {'t': '· Breakage（未兌換率）依歷史核銷率估計，降低負債高估', 'size': 13.5, 'color': WHITE, 'sa': 5},
    {'t': '· 商家/合作方分攤點數成本 → 北捷點數負債壓力可控', 'size': 13.5, 'color': WHITE, 'sa': 5},
    {'t': '· 研究底稿：OPENPOINT / LINE POINTS / 星巴克 / Marriott 案例分析', 'size': 12, 'color': MINT},
])
src = card(s, 8.75, 3.85, 4.0, 2.8, fill=PAPER, line_w=2.5, shadow=0.1)
txt(s, 9.05, 4.1, 3.5, 2.4, [
    {'t': '金流來源', 'size': 15, 'bold': True, 'color': TEAL, 'sa': 6},
    {'t': '· 商家行銷方案與導客抽成', 'size': 13, 'sa': 4},
    {'t': '· 站內租賃點位加值（招商）', 'size': 13, 'sa': 4},
    {'t': '· 健康共享合作案（市府）', 'size': 13, 'sa': 4},
    {'t': '· ESG 數據服務（後期）', 'size': 13},
])
footer(s, 11)

# ============ SLIDE 12 — 三階段藍圖 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 12, '三階段導入藍圖與 KPI')
phases = [
    ('決賽階段', '可操作概念原型\n+ 北捷 Dashboard 展示\n（不涉及深度串接）', MINT),
    ('試辦階段', '2–3 個站點生活圈\n串既有捷運點商家\n出站推播 + 換券核銷', SUN),
    ('擴大階段', '常客優惠轉點活化\n場館/觀光/共享運具\n悠遊付商店網絡', COIN),
]
for i, (t, d, cl) in enumerate(phases):
    x = 0.55 + i * 4.18
    c = card(s, x, 1.55, 3.85, 2.6, fill=PAPER, line_w=2.5, shadow=0.1)
    chip(s, x + 0.25, 1.8, 1.75, 0.55, t, fill=cl, size=14)
    txt(s, x + 0.28, 2.55, 3.3, 1.5, [{'t': line, 'size': 13, 'ls': 1.3, 'sa': 3} for line in d.split('\n')])
    if i < 2:
        txt(s, x + 3.85, 2.5, 0.4, 0.8, [{'t': '→', 'size': 18, 'bold': True, 'align': PP_ALIGN.CENTER}])
kpi = card(s, 0.55, 4.5, 12.2, 2.2, fill=TEAL, line_w=3, shadow=0.1)
txt(s, 0.95, 4.72, 11.4, 0.5, [{'t': '成效 KPI（試辦階段即可量測）', 'size': 15, 'bold': True, 'color': SUN}])
kpis = ['推播開啟率', '券核銷率', '點數消耗率', '估算導客量', '月減碳量 (kg)', '每導客成本']
for i, k in enumerate(kpis):
    x = 0.95 + (i % 6) * 1.95
    chip(s, x, 5.35, 1.8, 0.75, k, fill=PAPER, size=12.5)
txt(s, 0.95, 6.25, 11.4, 0.4, [{'t': '金融 ESG 聯名列為後期（法規與談判週期長）；團隊技術權利義務於試辦階段另行議定', 'size': 11.5, 'color': MINT}])
footer(s, 12)

# ============ SLIDE 13 — 誠實標註 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 13, '誠實標註：哪些是真的、哪些待確認', tag='HONESTY')
real = card(s, 0.55, 1.55, 6.0, 4.9, fill=PAPER, line_w=2.5, shadow=0.1)
txt(s, 0.9, 1.8, 5.4, 4.3, [
    {'t': '✓ 有依據的', 'size': 17, 'bold': True, 'color': LEAF, 'sa': 8},
    {'t': '· 北捷 114 年日運量 207.4 萬（官方統計）', 'size': 13, 'sa': 5},
    {'t': '· 站內租賃廠商皆為公開可查的真實案例', 'size': 13, 'sa': 5},
    {'t': '· U-Sport / U-Walk / YouBike 皆為既有系統', 'size': 13, 'sa': 5},
    {'t': '· 推播時段設計參考臺北市資料大平臺「各站分時進出量」開放資料', 'size': 13, 'sa': 5, 'ls': 1.3},
    {'t': '· 原型為可操作前端，決賽現場可離線執行', 'size': 13},
])
todo = card(s, 6.75, 1.55, 6.0, 4.9, fill=PAPER, line_w=2.5, shadow=0.1)
txt(s, 7.1, 1.8, 5.4, 4.3, [
    {'t': '△ 待北捷確認的', 'size': 17, 'bold': True, 'color': COIN, 'sa': 8},
    {'t': '· 碳係數 104 / 75 / 29 g／人km（口述參考值）', 'size': 13, 'sa': 5},
    {'t': '· 月活會員數（口述 20–30 萬）', 'size': 13, 'sa': 5},
    {'t': '· 首年財務試算之合作金流假設', 'size': 13, 'sa': 5},
    {'t': '· Demo 內點數 / 商家 / 儀表板數值為模擬情境', 'size': 13, 'sa': 5},
    {'t': '· 常客優惠轉捷運點之政策時程', 'size': 13},
])
b = card(s, 0.55, 6.5, 12.2, 0.5, fill=INK, line_w=2, shadow=0.06)
b.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(b.text_frame, [{'t': '風險透明化是可行性的一部分 — 我們把假設攤在桌上，導入時逐項替換為正式值', 'size': 13, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER, 'sa': 0}])
footer(s, 13)

# ============ SLIDE 14 — 團隊 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 14, '團隊：懂技術、懂用戶、也懂政府')
team = [
    ('嚴世紀', 'Numbers Protocol', '專案總指揮、系統架構、AI 推薦引擎、碳數據可信度驗證', TEAL),
    ('鍾天傑', 'LINE Pay', '點數經濟模型、加成機制設計、前端 UX 與社群激勵', COIN),
    ('嚴世裕', '台北市政府地政局', 'ESG 政策對接、開放資料整合、合規審查', LEAF),
]
for i, (name, org, role, cl) in enumerate(team):
    x = 0.55 + i * 4.18
    c = card(s, x, 1.6, 3.85, 3.1, fill=PAPER, line_w=2.5, shadow=0.1)
    txt(s, x + 0.3, 1.9, 3.2, 0.6, [{'t': name, 'size': 22, 'bold': True}])
    chip(s, x + 0.3, 2.6, 2.4, 0.55, org, fill=cl, tcolor=(WHITE if cl != COIN else INK), size=12.5)
    txt(s, x + 0.3, 3.35, 3.25, 1.2, [{'t': role, 'size': 12.5, 'ls': 1.35}])
tech = card(s, 0.55, 5.1, 12.2, 1.55, fill=DEEP, line_w=3, shadow=0.1)
txt(s, 0.95, 5.32, 11.4, 1.15, [
    {'t': '私部門技術力 × 點數經濟實戰 × 公部門政策力', 'size': 17, 'bold': True, 'color': SUN, 'sa': 4},
    {'t': '技術支持：Omni Edge（AI 智能體平台）· Numbers Protocol（Provenance infrastructure for humans & AI）', 'size': 12.5, 'color': MINT},
])
footer(s, 14)

# ============ SLIDE 15 — 收尾 ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
big = card(s, 0.55, 0.7, 12.2, 2.5, fill=TEAL, line_w=3, shadow=0.12)
big.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(big.text_frame, [
    {'t': '每一次通勤，都是可驗證的減碳資產', 'size': 32, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER, 'sa': 6},
    {'t': '從碳客變捷客 — 讓每一趟捷運，都成為省錢與減碳的證明', 'size': 15, 'color': MINT, 'align': PP_ALIGN.CENTER, 'sa': 0},
])
who = [('使用者', '省錢、減碳有感'), ('商家', '精準導客、成效可量化'), ('北捷', '點數活化、非票務收入'), ('城市', '可稽核的減碳貢獻')]
for i, (t, d) in enumerate(who):
    x = 0.55 + i * 3.13
    c = card(s, x, 3.5, 2.85, 1.35, fill=PAPER, line_w=2.5, shadow=0.08)
    txt(s, x + 0.22, 3.68, 2.4, 0.45, [{'t': t, 'size': 15, 'bold': True, 'color': TEAL}])
    txt(s, x + 0.22, 4.18, 2.45, 0.55, [{'t': d, 'size': 12.5}])
qrc = card(s, 0.55, 5.15, 2.0, 2.0, fill=PAPER, line_w=2.5, shadow=0.09, radius=0.1)
s.shapes.add_picture(f'{A}/qr-demo.png', Inches(0.67), Inches(5.27), width=Inches(1.76))
txt(s, 2.75, 5.35, 6.6, 1.6, [
    {'t': '互動原型（免安裝・可離線）', 'size': 14, 'bold': True, 'sa': 3},
    {'t': 'https://from-carbon-metro-vote.pages.dev/demo/', 'size': 14, 'bold': True, 'color': TEAL, 'font': 'Consolas', 'sa': 6},
    {'t': '線上人氣票選進行中：台北捷運Go App → 黑客松人氣票選 → 投給【從碳客變捷客】', 'size': 12.5, 'color': GREY},
])
fin = card(s, 9.6, 5.35, 3.2, 1.35, fill=SUN, line_w=2.5, shadow=0.09)
fin.rotation = -1.5
fin.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(fin.text_frame, [
    {'t': '黑捷客出發，創意進站', 'size': 14, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 2},
    {'t': '決賽我們全力到站！', 'size': 14, 'bold': True, 'align': PP_ALIGN.CENTER, 'sa': 0},
])

# ============ SLIDE 16 — 附錄 QA ============
s = prs.slides.add_slide(BLANK); slide_bg(s)
header(s, 'Q', '附錄：評審 QA 備彈')
qa = [
    ('點數成本誰出？和「商家出資點數池」的隊伍差在哪？',
     '我們對接北捷「既有招商業務」的加值：商家買的是行銷方案與導客成效，點數是方案內誘因；不需說服商家接受全新點數概念，北捷點數負債由分攤機制控管。'),
    ('碳係數怎麼驗證？被挑戰數字不準怎麼辦？',
     '係數為可配置參數（Demo 即可調），採北捷口述參考值並誠實標註；正式導入即替換為北捷/環境部正式係數，架構不變。每筆減碳紀錄附 Numbers 資料指紋，第三方可稽核。'),
    ('點數會計怎麼處理？',
     'IFRS 15：發點列遞延收入（合約負債），核銷/到期認列；breakage 依歷史核銷率估計。已完成 OPENPOINT、LINE POINTS、星巴克、Marriott 等案例研究底稿。'),
    ('和 Go App 的整合成本？',
     '定位為 Go App 加值模組：輕量 API + 推播 + 既有 Go!Map 資料，不動售票核心、不需新硬體；試辦僅需 2–3 站生活圈與既有捷運點商家。'),
]
y = 1.5
for q, a in qa:
    c = card(s, 0.55, y, 12.2, 1.28, fill=PAPER, line_w=2, shadow=0.06)
    txt(s, 0.9, y + 0.1, 11.6, 0.4, [{'t': 'Q：' + q, 'size': 13.5, 'bold': True, 'color': TEAL}])
    txt(s, 0.9, y + 0.52, 11.6, 0.7, [{'t': 'A：' + a, 'size': 11.5, 'ls': 1.2}])
    y += 1.42

out = 'personal-auto/final-deck/從碳客變捷客_決賽簡報_v2.pptx'
prs.save(out)
print('saved', out, '| slides:', len(prs.slides.__iter__.__self__._sldIdLst))
